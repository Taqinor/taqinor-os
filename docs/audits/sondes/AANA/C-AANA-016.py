"""Sonde AANA, rejouable par scripts/sonde.py : transaction annulée. Jamais contre la prod.
Rejoue C-AANA-016 (audit analyse du 05/10, §6) : rapports planifiés / lien public / export de définition sortaient prix_achat, valeur_achat.
"""
SONDE = {
    'constat': 'C-AANA-016',
    'sha': '5908d43bb',
    'attendu': "défaut absent : rendu planifié et lien public = en-tête ['nom'] ; export CSV d'un lecteur sans prix_achat_voir = ['nom'] sans 300.00",
}


def sonde(ctx):
    import csv
    import io
    from decimal import Decimal
    from django.contrib.auth import get_user_model
    from openpyxl import load_workbook
    from rest_framework.test import APIClient
    from apps.reporting import scheduled_reports
    from apps.reporting.diffusion_views import lien_rapport
    from apps.reporting.models import RapportDefinition, SavedReport
    from apps.roles.models import Role
    from apps.stock.models import Produit
    from authentication.models import Company
    from core.models import SavedQuery
    U = get_user_model()
    co = Company.objects.create(nom='Sonde AANA-016', slug='sonde-aana-016')
    admin = U.objects.create_user(username='sonde-aana-016-adm', password='x', role_legacy='admin', company=co)
    Produit.objects.create(company=co, nom='Panneau', sku='SONDE-AANA-016', prix_vente=Decimal('150'), prix_achat=Decimal('100'), quantite_stock=3)

    def entetes(contenu):
        return [c.value for c in next(load_workbook(io.BytesIO(contenu)).active.iter_rows(max_row=1))]
    q = SavedQuery.objects.create(company=co, owner=admin, titre='Catalogue achat', dataset='stock_produits', spec={'select': ['nom', 'prix_achat']})
    rap = SavedReport.objects.create(company=co, owner=admin, name='Catalogue', target_kind=SavedReport.TargetKind.QUERY, cible_id=q.pk, recipients='externe@example.invalid')
    contenu, _t, _n, _m = scheduled_reports.rendre_rapport(rap)
    h_plan = entetes(contenu)
    h_pub = entetes(APIClient().get(lien_rapport(rap), HTTP_HOST='localhost').content)
    d = RapportDefinition.objects.create(company=co, owner=admin, titre='Valeur stock', dataset='stock_produits', spec={'select': ['nom', 'valeur_achat']}, partage=RapportDefinition.Partage.SOCIETE)
    role = Role.objects.create(company=co, nom='Lecteur stock', permissions=['stock_voir', 'reporting_export'])
    lec = U.objects.create_user(username='sonde-aana-016-lec', password='x', company=co, role=role)
    api = APIClient()
    api.force_authenticate(lec)
    resp = api.get('/api/django/reporting/rapport-definitions/%s/export/?format=csv' % d.pk, HTTP_HOST='localhost')
    lignes = list(csv.reader(io.StringIO(resp.content.decode('utf-8-sig'))))
    fuite = '300.00' in {c for l in lignes for c in l}
    ok = h_plan == ['nom'] and h_pub == ['nom'] and resp.status_code == 200 and lignes[0] == ['nom'] and not fuite and not lec.can_view_buy_prices
    print('planifié=%s public=%s export=%s %s fuite300=%s' % (h_plan, h_pub, resp.status_code, lignes[0] if lignes else None, fuite))
    return {'repro': ok, 'planifie': h_plan, 'public': h_pub, 'export_entete': lignes[0] if lignes else None, 'fuite': fuite}
