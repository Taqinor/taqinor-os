"""ACRM4 (C-ACRM-002) — la règle UNIQUE du masquage PII sur toutes les
sorties de ``crm/views.py``.

Un rôle = ``COMMERCIAL_PERMISSIONS`` − ``client_pii_voir`` ne reçoit le numéro
témoin ``+212661909901`` dans AUCUN corps : historique, épingler, doublons
(fiche, pré-création, atelier), rapprochement client, autocomplete clients,
convertir-client, exports xlsx clients et leads (colonnes PII retirées),
partage WhatsApp de devis (403 ``droit_manquant``, comme ``resume_associe``).
Un Commercial complet obtient les réponses d'avant (numéro présent).

Sonde V_VA LVIEW1-2 : avant, les neuf chemins fuyaient. Aucun mock d'une
source interne (seul le fournisseur OCR EXTERNE est remplacé pour le scan de
carte de visite).
"""
import io
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import AccessToken

from authentication.models import Company

from apps.crm.models import Client, Lead, LeadActivity
from apps.roles.models import Role
from apps.roles.permissions_registre import COMMERCIAL_PERMISSIONS

User = get_user_model()
TEMOIN = '+212661909901'
FRAGMENT = '661909901'
LEADS = '/api/django/crm/leads/'
CLIENTS = '/api/django/crm/clients/'


def _cellules(response):
    classeur = load_workbook(io.BytesIO(response.content))
    return ' '.join(str(v) for row in classeur.active.iter_rows(
        values_only=True) for v in row if v is not None)


class PiiVuesTests(TestCase):

    def setUp(self):
        self.company = Company.objects.create(
            nom='ACRM4 Solaire', slug='acrm4-pii')
        complet = Role.objects.create(
            company=self.company, nom='Commercial',
            permissions=list(COMMERCIAL_PERMISSIONS), est_systeme=True)
        sans_pii = Role.objects.create(
            company=self.company, nom='Commercial sans PII',
            permissions=[p for p in COMMERCIAL_PERMISSIONS
                         if p != 'client_pii_voir'])
        self.complet = User.objects.create_user(
            username='acrm4-complet', password='x', company=self.company,
            role=complet)
        self.masque = User.objects.create_user(
            username='acrm4-masque', password='x', company=self.company,
            role=sans_pii)
        self.client_c = Client.objects.create(
            company=self.company, nom='Temoin', prenom='Client',
            telephone=TEMOIN, email='temoin-acrm4@example.com')
        self.lead = Lead.objects.create(
            company=self.company, nom='Temoin', prenom='Un',
            telephone=TEMOIN, email='temoin-acrm4@example.com',
            owner=self.masque, client=self.client_c)
        self.lead2 = Lead.objects.create(
            company=self.company, nom='Temoin', prenom='Deux',
            telephone=TEMOIN, owner=self.masque)
        # Les deux comptes voient les mêmes leads : le second compte possède
        # un jumeau de chaque lead pour rester dans sa portée.
        self.lead_c = Lead.objects.create(
            company=self.company, nom='Temoin', prenom='Trois',
            telephone=TEMOIN, email='temoin-acrm4@example.com',
            owner=self.complet, client=self.client_c)
        self.lead_c2 = Lead.objects.create(
            company=self.company, nom='Temoin', prenom='Quatre',
            telephone=TEMOIN, owner=self.complet)
        self.modif = LeadActivity.objects.create(
            company=self.company, lead=self.lead,
            kind=LeadActivity.Kind.MODIFICATION, field='telephone',
            old_value='+212600000000', new_value=TEMOIN)
        self.modif_c = LeadActivity.objects.create(
            company=self.company, lead=self.lead_c,
            kind=LeadActivity.Kind.MODIFICATION, field='telephone',
            old_value='+212600000000', new_value=TEMOIN)

    def _api(self, user):
        api = APIClient()
        api.credentials(HTTP_AUTHORIZATION=(
            f'Bearer {AccessToken.for_user(user)}'))
        return api

    def _corps(self, user, lead, modif):
        """Les réponses JSON des chemins lus (statut, texte)."""
        api = self._api(user)
        return {
            'historique': api.get(f'{LEADS}{lead.pk}/historique/'),
            'epingler': api.post(
                f'{LEADS}{lead.pk}/activites/{modif.pk}/epingler/'),
            'duplicates': api.get(f'{LEADS}{lead.pk}/duplicates/'),
            'check_duplicates': api.get(
                f'{LEADS}check-duplicates/', {'telephone': TEMOIN}),
            'doublons': api.get(f'{LEADS}doublons/'),
            'client_match': api.get(f'{LEADS}{lead.pk}/client-match/'),
            'clients_search': api.get(f'{CLIENTS}search/', {'q': 'Temoin'}),
            'convertir': api.post(
                f'{LEADS}{lead.pk}/convertir-client/',
                {'mode': 'lier', 'client_id': self.client_c.pk},
                format='json'),
        }

    def test_masque_sur_chaque_chemin_json(self):
        for chemin, resp in self._corps(
                self.masque, self.lead, self.modif).items():
            self.assertEqual(resp.status_code, 200, (chemin, resp.content))
            self.assertNotIn(FRAGMENT, resp.content.decode(), chemin)

    def test_convertir_client_lie_toujours(self):
        self._api(self.masque).post(
            f'{LEADS}{self.lead2.pk}/convertir-client/',
            {'mode': 'lier', 'client_id': self.client_c.pk}, format='json')
        self.lead2.refresh_from_db()
        self.assertEqual(self.lead2.client_id, self.client_c.pk)

    def test_exports_sans_colonnes_pii(self):
        api = self._api(self.masque)
        clients = api.post(f'{CLIENTS}export-xlsx/', {}, format='json')
        leads = api.post(f'{LEADS}export-xlsx/',
                         {'ids': [self.lead.pk, self.lead2.pk]},
                         format='json')
        for nom, resp in (('clients', clients), ('leads', leads)):
            self.assertEqual(resp.status_code, 200, nom)
            texte = _cellules(resp)
            self.assertIn('Temoin', texte, nom)
            self.assertNotIn(FRAGMENT, texte, nom)
            self.assertNotIn('Téléphone', texte, nom)

    def test_whatsapp_devis_refuse(self):
        api = self._api(self.masque)
        avant = LeadActivity.objects.filter(lead=self.lead).count()
        for chemin in ('whatsapp-devis-apercu', 'whatsapp-devis'):
            resp = api.post(f'{LEADS}{self.lead.pk}/{chemin}/', {},
                            format='json')
            self.assertEqual(resp.status_code, 403, chemin)
            self.assertEqual(resp.data.get('code'), 'droit_manquant', chemin)
            self.assertNotIn(FRAGMENT, resp.content.decode(), chemin)
        self.assertEqual(
            LeadActivity.objects.filter(lead=self.lead).count(), avant)

    def test_scan_carte_doublons_masques(self):
        self._scan_carte(self.masque, attendu_present=False)

    def test_commercial_complet_inchange(self):
        corps = self._corps(self.complet, self.lead_c, self.modif_c)
        for chemin in ('historique', 'epingler', 'duplicates',
                       'check_duplicates', 'doublons', 'client_match',
                       'clients_search', 'convertir'):
            self.assertEqual(corps[chemin].status_code, 200, chemin)
            self.assertIn(FRAGMENT, corps[chemin].content.decode(), chemin)
        api = self._api(self.complet)
        resp = api.post(f'{CLIENTS}export-xlsx/', {}, format='json')
        self.assertIn(FRAGMENT, _cellules(resp))
        resp = api.post(f'{LEADS}export-xlsx/', {'ids': [self.lead_c.pk]},
                        format='json')
        self.assertIn(FRAGMENT, _cellules(resp))
        self._scan_carte(self.complet, attendu_present=True)

    def _scan_carte(self, user, *, attendu_present):
        from django.core.files.uploadedfile import SimpleUploadedFile

        class _Resultat:
            configured = True
            data = {'nom': 'Temoin', 'telephone': TEMOIN}

        png = b'\x89PNG\r\n\x1a\n' + b'\x00' * 32
        with patch('core.ai.services.extract_document',
                   return_value=_Resultat()):
            resp = self._api(user).post(
                f'{LEADS}scan-carte/',
                {'file': SimpleUploadedFile('carte.png', png,
                                            content_type='image/png')},
                format='multipart')
        self.assertEqual(resp.status_code, 200, resp.content)
        doublons = str(resp.data['doublons'])
        self.assertTrue(resp.data['doublons'])
        self.assertEqual(FRAGMENT in doublons, attendu_present)
