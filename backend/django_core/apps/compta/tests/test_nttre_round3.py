"""Tests NTTRE round 3 — cockpit cash, exports et états imprimables.

NTTRE17 : bloc « Cash aujourd'hui » publié DANS ``etats/position-tresorerie/``
(solde consolidé du jour, delta vs la veille, 3 prochaines échéances) — la
carte du cockpit n'émet aucune requête supplémentaire.
NTTRE19 : export .xlsx du prévisionnel 13 semaines (colonnes semaine + ligne
« Solde projeté »), identique aux chiffres de l'écran.
"""
import io
from datetime import date, timedelta
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import AccessToken

from authentication.models import Company

from apps.compta import selectors, services
from apps.compta.models import (
    CompteTresorerie, Effet, LignePrevisionnelTresorerie, PaymentRun)

User = get_user_model()

JOUR = date(2026, 3, 10)


def make_company(slug, nom=None):
    company, _ = Company.objects.get_or_create(
        slug=slug, defaults={'nom': nom or slug})
    return company


def auth(user):
    api = APIClient()
    api.credentials(HTTP_AUTHORIZATION=f'Bearer {AccessToken.for_user(user)}')
    return api


def ecriture_banque(company, montant, jour):
    """Encaissement en banque : débit 5141 / crédit 3421 (écriture équilibrée)."""
    from apps.compta.models import Journal

    journal = services._journal(company, Journal.Type.BANQUE)
    return services.creer_ecriture(
        company, journal, jour, 'Encaissement test', [
            {'compte': services.get_compte(company, '5141'),
             'debit': Decimal(montant), 'credit': Decimal('0')},
            {'compte': services.get_compte(company, '3421'),
             'debit': Decimal('0'), 'credit': Decimal(montant)},
        ])


class CashAujourdhuiTests(TestCase):
    """NTTRE17 — solde du jour, delta vs la veille, 3 prochaines échéances."""

    def setUp(self):
        self.co = make_company('nttre17', 'NTTRE17 Co')
        services.seed_plan_comptable(self.co)
        services.seed_journaux(self.co)
        self.banque = CompteTresorerie.objects.create(
            company=self.co, type_compte=CompteTresorerie.Type.BANQUE,
            libelle='BMCE', solde_initial=Decimal('1000'),
            compte_comptable=services.get_compte(self.co, '5141'))

    def test_delta_veille_isole_les_mouvements_du_jour(self):
        ecriture_banque(self.co, '300', JOUR)
        bloc = selectors.cash_aujourdhui(self.co, aujourd_hui=JOUR)
        self.assertEqual(bloc['total'], Decimal('1300'))
        self.assertEqual(bloc['total_veille'], Decimal('1000'))
        self.assertEqual(bloc['delta_veille'], Decimal('300'))

    def test_trois_prochaines_echeances_effets_et_campagnes(self):
        # Effet à recevoir (+), effet à payer (−) et campagne de règlement (−).
        Effet.objects.create(
            company=self.co, sens=Effet.Sens.RECEVOIR,
            type_effet=Effet.TypeEffet.CHEQUE, numero='CHQ-1',
            montant=Decimal('500'), date_emission=JOUR,
            date_echeance=JOUR + timedelta(days=3))
        Effet.objects.create(
            company=self.co, sens=Effet.Sens.PAYER,
            type_effet=Effet.TypeEffet.TRAITE, numero='LCN-9',
            montant=Decimal('700'), date_emission=JOUR,
            date_echeance=JOUR + timedelta(days=5))
        # Effet DÉJÀ encaissé : hors échéancier.
        Effet.objects.create(
            company=self.co, sens=Effet.Sens.RECEVOIR,
            type_effet=Effet.TypeEffet.CHEQUE, numero='CHQ-SOLDE',
            montant=Decimal('999'), date_emission=JOUR,
            date_echeance=JOUR + timedelta(days=1),
            statut=Effet.Statut.ENCAISSE)
        PaymentRun.objects.create(
            company=self.co, reference='RUN-1',
            compte_tresorerie=self.banque,
            date_paiement=JOUR + timedelta(days=2),
            total=Decimal('200'))

        bloc = selectors.cash_aujourdhui(self.co, aujourd_hui=JOUR)
        echeances = bloc['prochaines_echeances']
        self.assertEqual(len(echeances), 3)
        # Chronologique : campagne J+2, effet à recevoir J+3, effet à payer J+5.
        self.assertEqual(
            [e['libelle'] for e in echeances], ['RUN-1', 'CHQ-1', 'LCN-9'])
        self.assertEqual(echeances[0]['montant'], Decimal('-200'))
        self.assertEqual(echeances[1]['montant'], Decimal('500'))
        self.assertEqual(echeances[2]['montant'], Decimal('-700'))

    def test_maximum_trois_echeances(self):
        for i in range(6):
            Effet.objects.create(
                company=self.co, sens=Effet.Sens.RECEVOIR,
                type_effet=Effet.TypeEffet.CHEQUE, numero=f'CHQ-{i}',
                montant=Decimal('100'), date_emission=JOUR,
                date_echeance=JOUR + timedelta(days=i + 1))
        bloc = selectors.cash_aujourdhui(self.co, aujourd_hui=JOUR)
        self.assertEqual(len(bloc['prochaines_echeances']), 3)

    def test_isolation_societe(self):
        autre = make_company('nttre17-b', 'NTTRE17 B')
        services.seed_plan_comptable(autre)
        CompteTresorerie.objects.create(
            company=autre, type_compte=CompteTresorerie.Type.BANQUE,
            libelle='Autre banque', solde_initial=Decimal('9999'),
            compte_comptable=services.get_compte(autre, '5141'))
        bloc = selectors.cash_aujourdhui(self.co, aujourd_hui=JOUR)
        self.assertEqual(bloc['total'], Decimal('1000'))

    def test_endpoint_position_publie_le_bloc_sans_appel_supplementaire(self):
        user = User.objects.create_user(
            username='nttre17-user', password='x', company=self.co,
            role_legacy='responsable')
        resp = auth(user).get('/api/django/compta/etats/position-tresorerie/')
        self.assertEqual(resp.status_code, 200)
        self.assertIn('cash_du_jour', resp.data)
        bloc = resp.data['cash_du_jour']
        for cle in ('total', 'total_veille', 'delta_veille',
                    'prochaines_echeances'):
            self.assertIn(cle, bloc)


class PrevisionnelXlsxTests(TestCase):
    """NTTRE19 — classeur banquier : 13 colonnes semaine + solde projeté."""

    def setUp(self):
        self.co = make_company('nttre19', 'NTTRE19 Co')
        services.seed_plan_comptable(self.co)
        services.seed_journaux(self.co)
        CompteTresorerie.objects.create(
            company=self.co, type_compte=CompteTresorerie.Type.BANQUE,
            libelle='BMCE', solde_initial=Decimal('1000'),
            compte_comptable=services.get_compte(self.co, '5141'))
        self.user = User.objects.create_user(
            username='nttre19-user', password='x', company=self.co,
            role_legacy='responsable')
        self.api = auth(self.user)

    def _feuille(self, contenu):
        from openpyxl import load_workbook
        return load_workbook(io.BytesIO(contenu)).active

    def test_treize_colonnes_semaine_et_ligne_solde_projete(self):
        resp = self.api.get(
            '/api/django/compta/etats/previsionnel-tresorerie/?export=xlsx')
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheet', resp['Content-Type'])
        ws = self._feuille(resp.content)
        # 1 colonne de libellé + 13 colonnes semaine.
        self.assertEqual(ws.max_column, 14)
        libelles = [ws.cell(row=r, column=1).value
                    for r in range(2, ws.max_row + 1)]
        self.assertIn('Solde projeté', libelles)

    def test_chiffres_identiques_a_ceux_de_l_ecran(self):
        # Le prévisionnel démarre au LUNDI de la semaine courante : on cale la
        # ligne prévue sur cette même semaine (aucune horloge figée requise).
        aujourdhui = timezone.localdate()
        lundi = aujourdhui - timedelta(days=aujourdhui.weekday())
        LignePrevisionnelTresorerie.objects.create(
            company=self.co, libelle='Subvention',
            date_prevue=lundi + timedelta(days=2), montant=Decimal('750'))
        ecran = self.api.get(
            '/api/django/compta/etats/previsionnel-tresorerie/')
        classeur = self.api.get(
            '/api/django/compta/etats/previsionnel-tresorerie/?export=xlsx')
        self.assertEqual(ecran.status_code, 200)
        self.assertEqual(classeur.status_code, 200)
        ws = self._feuille(classeur.content)
        lignes = {ws.cell(row=r, column=1).value: [
            ws.cell(row=r, column=c).value
            for c in range(2, ws.max_column + 1)]
            for r in range(2, ws.max_row + 1)}
        soldes_ecran = [float(s['solde_fin'])
                        for s in ecran.data['semaines']]
        self.assertEqual(lignes['Solde projeté'], soldes_ecran)
        self.assertEqual(lignes['Encaissements'][0], 750.0)

    def test_nb_semaines_personnalise_change_le_nombre_de_colonnes(self):
        resp = self.api.get(
            '/api/django/compta/etats/previsionnel-tresorerie/'
            '?export=xlsx&nb_semaines=4')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(self._feuille(resp.content).max_column, 5)


class QualiteRapprochementsTests(TestCase):
    """NTTRE20 — lignes restées non pointées à la clôture, mois par mois."""

    def setUp(self):
        self.co = make_company('nttre20', 'NTTRE20 Co')
        services.seed_plan_comptable(self.co)
        services.seed_journaux(self.co)
        self.banque = CompteTresorerie.objects.create(
            company=self.co, type_compte=CompteTresorerie.Type.BANQUE,
            libelle='BMCE', compte_comptable=services.get_compte(self.co, '5141'))

    def _rapprochement(self, mois, *, non_pointees=0, cloture=True):
        from apps.compta.models import RapprochementBancaire

        rap = services.creer_rapprochement(
            self.co, self.banque, date_debut=date(2026, mois, 1),
            date_fin=date(2026, mois, 28), solde_releve=Decimal('0'))
        for i in range(non_pointees):
            services.ajouter_ligne_releve(
                rap, date_operation=date(2026, mois, 5),
                libelle=f'Op {mois}-{i}', montant=Decimal('10'))
        if cloture:
            # Clôture POSÉE directement : le service refuse de clôturer tant
            # qu'un écart subsiste — on teste ici la LECTURE d'un historique
            # déjà clôturé (reprise d'antériorité, clôture forcée en base).
            rap.statut = RapprochementBancaire.Statut.RAPPROCHE
            rap.save(update_fields=['statut'])
        return rap

    def test_compte_exact_des_lignes_non_pointees_par_mois(self):
        self._rapprochement(2, non_pointees=3)
        self._rapprochement(3, non_pointees=1)
        data = selectors.qualite_rapprochements(
            self.co, nb_mois=12, today=date(2026, 12, 31))
        par_mois = {m['mois']: m for m in data['mois']}
        self.assertEqual(par_mois['2026-02']['lignes_non_pointees'], 3)
        self.assertEqual(par_mois['2026-02']['rapprochements'], 1)
        self.assertEqual(par_mois['2026-03']['lignes_non_pointees'], 1)
        self.assertEqual(data['total_lignes_non_pointees'], 4)

    def test_rapprochement_non_cloture_exclu(self):
        self._rapprochement(4, non_pointees=5, cloture=False)
        data = selectors.qualite_rapprochements(
            self.co, nb_mois=12, today=date(2026, 12, 31))
        self.assertEqual(data['total_lignes_non_pointees'], 0)
        par_mois = {m['mois']: m for m in data['mois']}
        self.assertEqual(par_mois['2026-04']['rapprochements'], 0)

    def test_fenetre_de_mois_continue_et_ordonnee(self):
        data = selectors.qualite_rapprochements(
            self.co, nb_mois=6, today=date(2026, 6, 15))
        self.assertEqual(
            [m['mois'] for m in data['mois']],
            ['2026-01', '2026-02', '2026-03', '2026-04', '2026-05', '2026-06'])

    def test_endpoint_scope_societe(self):
        self._rapprochement(2, non_pointees=2)
        autre = make_company('nttre20-b', 'NTTRE20 B')
        user_b = User.objects.create_user(
            username='nttre20-user-b', password='x', company=autre,
            role_legacy='responsable')
        resp = auth(user_b).get(
            '/api/django/compta/etats/qualite-rapprochements/')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['total_lignes_non_pointees'], 0)


class SituationEffetsTests(TestCase):
    """NTTRE22 — situation du portefeuille d'effets par statut/sens/tranche."""

    REFERENCE = date(2026, 5, 4)

    def setUp(self):
        self.co = make_company('nttre22', 'NTTRE22 Co')
        self.user = User.objects.create_user(
            username='nttre22-user', password='x', company=self.co,
            role_legacy='responsable')
        self.api = auth(self.user)

    def _effet(self, jours, montant, *, sens=Effet.Sens.RECEVOIR,
               statut=Effet.Statut.PORTEFEUILLE, numero=''):
        echeance = self.REFERENCE + timedelta(days=jours)
        return Effet.objects.create(
            company=self.co, sens=sens, statut=statut,
            type_effet=Effet.TypeEffet.CHEQUE,
            numero=numero or f'E{jours}',
            montant=Decimal(montant),
            # Émission toujours ANTÉRIEURE à l'échéance (invariant du modèle).
            date_emission=echeance - timedelta(days=30),
            date_echeance=echeance)

    def test_total_par_tranche_egale_la_somme_des_effets_qui_y_tombent(self):
        self._effet(5, '100')      # < 30 j
        self._effet(29, '50')      # < 30 j (borne haute exclue)
        self._effet(30, '200')     # 30-60 j (borne basse incluse)
        self._effet(75, '300')     # 60-90 j
        self._effet(120, '400')    # > 90 j
        data = selectors.situation_effets(
            self.co, date_reference=self.REFERENCE)
        par_code = {t['code']: t for t in data['tranches']}
        self.assertEqual(par_code['moins_30']['total'], Decimal('150'))
        self.assertEqual(par_code['moins_30']['nb'], 2)
        self.assertEqual(par_code['de_30_a_60']['total'], Decimal('200'))
        self.assertEqual(par_code['de_60_a_90']['total'], Decimal('300'))
        self.assertEqual(par_code['plus_90']['total'], Decimal('400'))
        self.assertEqual(data['total_general'], Decimal('1050'))
        # Chaque tranche redit exactement la somme des effets qu'elle liste.
        for tranche in data['tranches']:
            somme = sum((e['montant'] for e in tranche['effets']), Decimal('0'))
            self.assertEqual(tranche['total'], somme)

    def test_effet_en_retard_tombe_dans_la_premiere_tranche(self):
        self._effet(-10, '90')
        data = selectors.situation_effets(
            self.co, date_reference=self.REFERENCE)
        par_code = {t['code']: t for t in data['tranches']}
        self.assertEqual(par_code['moins_30']['total'], Decimal('90'))

    def test_groupes_par_sens_et_statut(self):
        self._effet(10, '100', sens=Effet.Sens.RECEVOIR,
                    statut=Effet.Statut.PORTEFEUILLE, numero='R1')
        self._effet(12, '60', sens=Effet.Sens.RECEVOIR,
                    statut=Effet.Statut.REMIS, numero='R2')
        self._effet(15, '80', sens=Effet.Sens.PAYER,
                    statut=Effet.Statut.PORTEFEUILLE, numero='P1')
        data = selectors.situation_effets(
            self.co, date_reference=self.REFERENCE)
        cles = {(g['sens'], g['statut']): g for g in data['groupes']}
        self.assertEqual(len(cles), 3)
        self.assertEqual(
            cles[('recevoir', 'portefeuille')]['total'], Decimal('100'))
        self.assertEqual(cles[('recevoir', 'remis')]['total'], Decimal('60'))
        self.assertEqual(
            cles[('payer', 'portefeuille')]['total'], Decimal('80'))

    def test_html_imprimable_reprend_les_tranches(self):
        from apps.compta.pdf_etats import render_situation_effets_html

        self._effet(5, '100', numero='CHQ-42')
        data = selectors.situation_effets(
            self.co, date_reference=self.REFERENCE)
        html = render_situation_effets_html(data, None, today=self.REFERENCE)
        self.assertIn('Situation des effets en portefeuille', html)
        self.assertIn('Moins de 30 jours', html)
        self.assertIn('CHQ-42', html)

    def test_endpoint_json_et_isolation_societe(self):
        self._effet(5, '100')
        resp = self.api.get(
            '/api/django/compta/etats/situation-effets/'
            f'?date={self.REFERENCE.isoformat()}')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['total_general'], Decimal('100'))

        autre = make_company('nttre22-b', 'NTTRE22 B')
        user_b = User.objects.create_user(
            username='nttre22-user-b', password='x', company=autre,
            role_legacy='responsable')
        resp_b = auth(user_b).get(
            '/api/django/compta/etats/situation-effets/')
        self.assertEqual(resp_b.status_code, 200)
        self.assertEqual(resp_b.data['nb_effets'], 0)

    def test_date_invalide_refusee_en_francais(self):
        resp = self.api.get(
            '/api/django/compta/etats/situation-effets/?date=pas-une-date')
        self.assertEqual(resp.status_code, 400)
        self.assertIn('date', resp.data['detail'])


class CertificatPouvoirBancaireTests(TestCase):
    """NTTRE23 — certificat PDF d'un pouvoir bancaire (mention « révoqué »)."""

    def setUp(self):
        self.co = make_company('nttre23', 'NTTRE23 Co')
        services.seed_plan_comptable(self.co)
        self.banque = CompteTresorerie.objects.create(
            company=self.co, type_compte=CompteTresorerie.Type.BANQUE,
            libelle='BMCE Agdal', banque='BMCE Bank',
            rib='011780000012345678901234',
            compte_comptable=services.get_compte(self.co, '5141'))
        self.user = User.objects.create_user(
            username='nttre23-user', password='x', company=self.co,
            role_legacy='responsable')
        self.api = auth(self.user)

    def _pouvoir(self, **kwargs):
        from apps.compta.models import PouvoirBancaire

        defaults = {
            'company': self.co, 'compte_tresorerie': self.banque,
            'titulaire_nom': 'Titulaire A', 'titulaire_cin': 'AB12345',
            'plafond_signature_seul': Decimal('100000'),
            'plafond_signature_conjointe': Decimal('500000'),
            'date_debut': date(2026, 1, 1), 'date_fin': date(2026, 12, 31),
        }
        defaults.update(kwargs)
        return PouvoirBancaire.objects.create(**defaults)

    def _html(self, pouvoir):
        from apps.compta.models import PouvoirBancaire
        from apps.compta.pdf_etats import render_certificat_pouvoir_html

        return render_certificat_pouvoir_html({
            'titulaire_nom': pouvoir.titulaire_nom,
            'titulaire_cin': pouvoir.titulaire_cin,
            'compte_libelle': pouvoir.compte_tresorerie.libelle,
            'compte_banque': pouvoir.compte_tresorerie.banque,
            'compte_rib': pouvoir.compte_tresorerie.rib,
            'plafond_signature_seul': pouvoir.plafond_signature_seul,
            'plafond_signature_conjointe': pouvoir.plafond_signature_conjointe,
            'date_debut': pouvoir.date_debut,
            'date_fin': pouvoir.date_fin,
            'statut': pouvoir.statut,
            'statut_libelle': pouvoir.get_statut_display(),
            'revoque': pouvoir.statut == PouvoirBancaire.Statut.REVOQUE,
        }, None, today=date(2026, 6, 1))

    def test_certificat_reprend_tous_les_champs_du_pouvoir_actif(self):
        html = self._html(self._pouvoir())
        self.assertIn('Certificat de pouvoir bancaire', html)
        self.assertIn('Titulaire A', html)
        self.assertIn('AB12345', html)
        self.assertIn('BMCE Agdal', html)
        self.assertIn('011780000012345678901234', html)
        self.assertIn('100 000,00', html)
        self.assertIn('500 000,00', html)
        self.assertIn('2026-01-01', html)
        self.assertIn('2026-12-31', html)
        self.assertNotIn('POUVOIR RÉVOQUÉ', html)

    def test_pouvoir_revoque_sort_barre_avec_la_mention(self):
        from apps.compta.models import PouvoirBancaire

        html = self._html(self._pouvoir(
            statut=PouvoirBancaire.Statut.REVOQUE))
        self.assertIn('POUVOIR RÉVOQUÉ', html)
        self.assertIn('line-through', html)
        self.assertIn('class="revoque"', html)

    def test_endpoint_pdf_ou_503(self):
        pouvoir = self._pouvoir()
        resp = self.api.get(
            f'/api/django/compta/pouvoirs-bancaires/{pouvoir.pk}/certificat/')
        self.assertIn(resp.status_code, (200, 503))
        if resp.status_code == 200:
            self.assertEqual(resp['Content-Type'], 'application/pdf')

    def test_certificat_autre_societe_introuvable(self):
        autre = make_company('nttre23-b', 'NTTRE23 B')
        services.seed_plan_comptable(autre)
        banque_b = CompteTresorerie.objects.create(
            company=autre, type_compte=CompteTresorerie.Type.BANQUE,
            libelle='Autre', compte_comptable=services.get_compte(autre, '5141'))
        pouvoir_b = self._pouvoir(company=autre, compte_tresorerie=banque_b,
                                  titulaire_nom='Titulaire B')
        resp = self.api.get(
            f'/api/django/compta/pouvoirs-bancaires/{pouvoir_b.pk}/certificat/')
        self.assertEqual(resp.status_code, 404)


class ApercuCampagnePaiementTests(TestCase):
    """NTTRE25 — sélection guidée + alerte de franchissement du seuil bas."""

    def setUp(self):
        from apps.stock.models import FactureFournisseur, Fournisseur

        self.co = make_company('nttre25', 'NTTRE25 Co')
        services.seed_plan_comptable(self.co)
        services.seed_journaux(self.co)
        self.banque = CompteTresorerie.objects.create(
            company=self.co, type_compte=CompteTresorerie.Type.BANQUE,
            libelle='BMCE', solde_initial=Decimal('10000'),
            seuil_alerte_bas=Decimal('5000'),
            compte_comptable=services.get_compte(self.co, '5141'))
        self.f1 = Fournisseur.objects.create(
            company=self.co, nom='Fournisseur A',
            rib='011780000012345678901234')
        self.f2 = Fournisseur.objects.create(
            company=self.co, nom='Fournisseur B',
            rib='011780000012345678905678')
        self.echeance = date(2026, 7, 15)
        FactureFournisseur.objects.create(
            company=self.co, reference='FF-A1', fournisseur=self.f1,
            date_echeance=self.echeance, montant_ht=Decimal('3000'),
            montant_tva=Decimal('0'), montant_ttc=Decimal('3000'))
        FactureFournisseur.objects.create(
            company=self.co, reference='FF-B1', fournisseur=self.f2,
            date_echeance=self.echeance + timedelta(days=30),
            montant_ht=Decimal('4000'), montant_tva=Decimal('0'),
            montant_ttc=Decimal('4000'))
        self.user = User.objects.create_user(
            username='nttre25-user', password='x', company=self.co,
            role_legacy='responsable')
        self.api = auth(self.user)

    def test_filtres_fournisseur_echeance_et_montant(self):
        tout = selectors.apercu_campagne_paiement(self.co)
        self.assertEqual(tout['nb_dettes'], 2)
        self.assertEqual(tout['total'], Decimal('7000'))

        par_fournisseur = selectors.apercu_campagne_paiement(
            self.co, fournisseur_id=self.f1.id)
        self.assertEqual(par_fournisseur['nb_dettes'], 1)
        self.assertEqual(par_fournisseur['total'], Decimal('3000'))

        par_echeance = selectors.apercu_campagne_paiement(
            self.co, date_limite=self.echeance)
        self.assertEqual(par_echeance['nb_dettes'], 1)

        par_montant = selectors.apercu_campagne_paiement(
            self.co, montant_max=Decimal('3500'))
        self.assertEqual(par_montant['nb_dettes'], 1)
        self.assertEqual(par_montant['total'], Decimal('3000'))

    def test_alerte_quand_le_compte_passerait_sous_son_seuil(self):
        # Solde 10 000, seuil 5 000, campagne 7 000 → projeté 3 000 < seuil.
        data = selectors.apercu_campagne_paiement(
            self.co, compte_tresorerie_id=self.banque.id)
        self.assertEqual(data['compte']['solde_actuel'], Decimal('10000'))
        self.assertEqual(data['compte']['solde_projete'], Decimal('3000'))
        self.assertIsNotNone(data['alerte_seuil'])
        self.assertEqual(data['alerte_seuil']['motif'], 'seuil_alerte_bas')

    def test_pas_d_alerte_quand_le_solde_reste_au_dessus(self):
        data = selectors.apercu_campagne_paiement(
            self.co, compte_tresorerie_id=self.banque.id,
            montant_max=Decimal('3500'))  # 3 000 seulement → projeté 7 000
        self.assertEqual(data['compte']['solde_projete'], Decimal('7000'))
        self.assertIsNone(data['alerte_seuil'])

    def test_compte_sans_seuil_ne_declenche_aucune_alerte(self):
        self.banque.seuil_alerte_bas = None
        self.banque.save(update_fields=['seuil_alerte_bas'])
        data = selectors.apercu_campagne_paiement(
            self.co, compte_tresorerie_id=self.banque.id)
        self.assertIsNone(data['alerte_seuil'])

    def test_proposer_applique_les_memes_filtres_que_l_apercu(self):
        run = services.creer_payment_run(
            self.co, date_paiement=self.echeance, mode_paiement='virement',
            compte_tresorerie=self.banque, reference='RUN-NTTRE25')
        ajoutees = services.proposer_lignes_payment_run(
            run, fournisseur_id=self.f1.id)
        self.assertEqual(len(ajoutees), 1)
        self.assertEqual(ajoutees[0].beneficiaire, 'Fournisseur A')
        run.refresh_from_db()
        self.assertEqual(run.total, Decimal('3000'))
        self.assertEqual(run.statut, run.Statut.BROUILLON)

    def test_endpoint_apercu_scope_societe_et_refuse_un_montant_non_numerique(self):
        resp = self.api.get(
            '/api/django/compta/payment-runs/apercu/'
            f'?compte={self.banque.id}')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['nb_dettes'], 2)

        mauvais = self.api.get(
            '/api/django/compta/payment-runs/apercu/?montant_max=beaucoup')
        self.assertEqual(mauvais.status_code, 400)
        self.assertIn('montant_max', mauvais.data)

        autre = make_company('nttre25-b', 'NTTRE25 B')
        user_b = User.objects.create_user(
            username='nttre25-user-b', password='x', company=autre,
            role_legacy='responsable')
        resp_b = auth(user_b).get('/api/django/compta/payment-runs/apercu/')
        self.assertEqual(resp_b.status_code, 200)
        self.assertEqual(resp_b.data['nb_dettes'], 0)


class ImportPlafondsPouvoirsTests(TestCase):
    """NTTRE37 — import CSV des plafonds : additif, jamais de création."""

    def setUp(self):
        self.co = make_company('nttre37', 'NTTRE37 Co')
        services.seed_plan_comptable(self.co)
        self.banque = CompteTresorerie.objects.create(
            company=self.co, type_compte=CompteTresorerie.Type.BANQUE,
            libelle='BMCE', compte_comptable=services.get_compte(self.co, '5141'))
        self.user = User.objects.create_user(
            username='nttre37-user', password='x', company=self.co,
            role_legacy='responsable')
        self.api = auth(self.user)

    def _pouvoir(self, nom, cin, seul=Decimal('0'), conjoint=Decimal('0')):
        from apps.compta.models import PouvoirBancaire

        return PouvoirBancaire.objects.create(
            company=self.co, compte_tresorerie=self.banque,
            titulaire_nom=nom, titulaire_cin=cin,
            plafond_signature_seul=seul,
            plafond_signature_conjointe=conjoint)

    @staticmethod
    def _csv(lignes):
        entete = 'cin;plafond_signature_seul;plafond_signature_conjointe'
        return ('\n'.join([entete] + lignes)).encode('utf-8')

    def test_cinquante_lignes_mettent_a_jour_les_plafonds(self):
        pouvoirs = [self._pouvoir(f'Titulaire {i}', f'CIN{i:03d}')
                    for i in range(50)]
        contenu = self._csv([f'CIN{i:03d};{1000 + i};{5000 + i}'
                             for i in range(50)])
        rapport = services.importer_plafonds_pouvoirs_csv(
            self.co, contenu, 'plafonds.csv', user=self.user)
        self.assertEqual(rapport['maj'], 50)
        self.assertEqual(rapport['crees'], 0)
        self.assertEqual(rapport['erreurs'], [])
        pouvoirs[0].refresh_from_db()
        self.assertEqual(pouvoirs[0].plafond_signature_seul, Decimal('1000.00'))
        self.assertEqual(
            pouvoirs[0].plafond_signature_conjointe, Decimal('5000.00'))
        pouvoirs[49].refresh_from_db()
        self.assertEqual(pouvoirs[49].plafond_signature_seul, Decimal('1049.00'))

    def test_cin_inconnue_rejetee_explicitement_sans_creation(self):
        from apps.compta.models import PouvoirBancaire

        self._pouvoir('Titulaire A', 'CIN001')
        contenu = self._csv(['CIN001;2000;8000', 'CIN999;3000;9000'])
        rapport = services.importer_plafonds_pouvoirs_csv(
            self.co, contenu, 'plafonds.csv', user=self.user)
        self.assertEqual(rapport['maj'], 1)
        self.assertEqual(len(rapport['erreurs']), 1)
        self.assertEqual(rapport['erreurs'][0]['ligne'], 2)
        self.assertIn('CIN999', rapport['erreurs'][0]['motif'])
        self.assertIn('introuvable', rapport['erreurs'][0]['motif'])
        # Aucun pouvoir créé par l'import.
        self.assertEqual(
            PouvoirBancaire.objects.filter(company=self.co).count(), 1)

    def test_remplissage_seul_par_defaut_ecrasement_sur_option(self):
        pouvoir = self._pouvoir(
            'Titulaire B', 'CIN002', seul=Decimal('100000'))
        contenu = self._csv(['CIN002;250000;'])
        rapport = services.importer_plafonds_pouvoirs_csv(
            self.co, contenu, 'plafonds.csv', user=self.user)
        pouvoir.refresh_from_db()
        # Plafond DÉJÀ non nul : non remplacé sans opt-in explicite.
        self.assertEqual(pouvoir.plafond_signature_seul, Decimal('100000.00'))
        self.assertEqual(len(rapport['refuses']), 1)

        rapport2 = services.importer_plafonds_pouvoirs_csv(
            self.co, contenu, 'plafonds.csv', user=self.user, ecraser=True)
        pouvoir.refresh_from_db()
        self.assertEqual(pouvoir.plafond_signature_seul, Decimal('250000.00'))
        self.assertEqual(len(rapport2['ecrasements']), 1)

    def test_montant_invalide_met_la_ligne_en_erreur(self):
        self._pouvoir('Titulaire C', 'CIN003')
        contenu = self._csv(['CIN003;beaucoup;1000'])
        rapport = services.importer_plafonds_pouvoirs_csv(
            self.co, contenu, 'plafonds.csv', user=self.user)
        self.assertEqual(rapport['maj'], 0)
        self.assertEqual(len(rapport['erreurs']), 1)
        self.assertIn('plafond_signature_seul', rapport['erreurs'][0]['motif'])

    def test_apercu_n_ecrit_rien(self):
        pouvoir = self._pouvoir('Titulaire D', 'CIN004')
        contenu = self._csv(['CIN004;7000;9000'])
        rapport = services.importer_plafonds_pouvoirs_csv(
            self.co, contenu, 'plafonds.csv', user=self.user, apercu=True)
        self.assertTrue(rapport['apercu'])
        self.assertEqual(rapport['maj'], 1)
        pouvoir.refresh_from_db()
        self.assertEqual(pouvoir.plafond_signature_seul, Decimal('0.00'))

    def test_le_lot_est_journalise_dans_importjob(self):
        from apps.dataimport.models import ImportJob

        self._pouvoir('Titulaire E', 'CIN005')
        contenu = self._csv(['CIN005;1500;2500', 'CIN888;1;2'])
        rapport = services.importer_plafonds_pouvoirs_csv(
            self.co, contenu, 'plafonds.csv', user=self.user)
        job = ImportJob.objects.get(pk=rapport['job_id'])
        self.assertEqual(job.company_id, self.co.id)
        self.assertEqual(job.target, 'plafonds_pouvoirs_bancaires')
        self.assertEqual(job.updated_count, 1)
        self.assertEqual(job.created_count, 0)
        self.assertEqual(job.error_count, 1)

    def test_endpoint_import_csv(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        pouvoir = self._pouvoir('Titulaire F', 'CIN006')
        fichier = SimpleUploadedFile(
            'plafonds.csv', self._csv(['CIN006;4200;']),
            content_type='text/csv')
        resp = self.api.post(
            '/api/django/compta/pouvoirs-bancaires/import-csv/',
            {'fichier': fichier}, format='multipart')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['maj'], 1)
        pouvoir.refresh_from_db()
        self.assertEqual(pouvoir.plafond_signature_seul, Decimal('4200.00'))

    def test_endpoint_refuse_un_fichier_manquant_en_nommant_le_champ(self):
        resp = self.api.post(
            '/api/django/compta/pouvoirs-bancaires/import-csv/',
            {}, format='multipart')
        self.assertEqual(resp.status_code, 400)
        self.assertIn('fichier', resp.data)

    def test_import_ne_voit_pas_les_pouvoirs_d_une_autre_societe(self):
        autre = make_company('nttre37-b', 'NTTRE37 B')
        services.seed_plan_comptable(autre)
        banque_b = CompteTresorerie.objects.create(
            company=autre, type_compte=CompteTresorerie.Type.BANQUE,
            libelle='Autre', compte_comptable=services.get_compte(autre, '5141'))
        from apps.compta.models import PouvoirBancaire
        pouvoir_b = PouvoirBancaire.objects.create(
            company=autre, compte_tresorerie=banque_b,
            titulaire_nom='Titulaire B', titulaire_cin='CIN007')
        rapport = services.importer_plafonds_pouvoirs_csv(
            self.co, self._csv(['CIN007;9999;']), 'plafonds.csv',
            user=self.user)
        self.assertEqual(rapport['maj'], 0)
        self.assertEqual(len(rapport['erreurs']), 1)
        pouvoir_b.refresh_from_db()
        self.assertEqual(pouvoir_b.plafond_signature_seul, Decimal('0.00'))
