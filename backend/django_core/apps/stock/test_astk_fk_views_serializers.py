"""ASTK7 (C-ASTK-001) — les sérialiseurs DÉCLARÉS dans apps/stock/views/*.py
et les ids bruts des services n'acceptent plus l'id d'une autre société.

Sondes d'origine : FOUR-1 (incident qualité sur le fournisseur de B → 201,
``fournisseur_nom`` = FOURNISSEUR-B-SECRET, puis la suppression du
fournisseur de B levait ``ProtectedError``) ; TEN-6(5) (PATCH consignation
vers un client de B → 200) ; TEN-6(8) (accord RFA sur un fournisseur de B →
201).

Désormais : ``CompanyScopedRelationsMixin`` sur chaque sérialiseur de ces
vues, ``creer_depot_consignation`` relit le client borné à la société, et
``valider_retour_scanne`` relit les casiers sources bornés à la société.
Réponse : 400 « objet inexistant », comme l'id 99999999.

Aucun mock : mixin core, vues et services réels.

Run :
    python manage.py test apps.stock.test_astk_fk_views_serializers -v 2
"""
import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import AccessToken

from apps.stock.models import (
    Categorie, DepotConsignation, EmplacementStock, Fournisseur,
    IncidentQualiteFournisseur, LigneRetourFournisseur, MouvementStock,
    Produit, RetourFournisseur,
)
from apps.stock.services_consignation import creer_depot_consignation

User = get_user_model()

BASE = '/api/django/stock'
JOUR = '2026-10-01'


def make_company(slug, nom):
    from authentication.models import Company
    company, _ = Company.objects.get_or_create(slug=slug, defaults={'nom': nom})
    return company


def auth_client(user):
    client = APIClient()
    client.credentials(
        HTTP_AUTHORIZATION=f'Bearer {AccessToken.for_user(user)}')
    return client


class FkViewsSerializersTests(TestCase):
    def setUp(self):
        from apps.crm.models import Client
        from apps.installations.models import BinLocation

        self.a = make_company('astk7-a', 'ASTK7 A')
        self.b = make_company('astk7-b', 'ASTK7 B')
        self.resp_a = User.objects.create_user(
            username='astk7_resp_a', password='x', role_legacy='responsable',
            company=self.a)
        self.api = auth_client(self.resp_a)

        # --- A -----------------------------------------------------------
        self.fa = Fournisseur.objects.create(company=self.a, nom='Fourn. A')
        self.pa = Produit.objects.create(
            company=self.a, nom='ProdA', sku='ASTK7-PA', fournisseur=self.fa,
            prix_achat=Decimal('10'), prix_vente=Decimal('20'),
            quantite_stock=50)
        self.client_a = Client.objects.create(
            company=self.a, nom='Client A ASTK7')
        emp_a = EmplacementStock.objects.create(
            company=self.a, nom='Dépôt A', is_principal=True)
        self.bin_a = BinLocation.objects.create(
            company=self.a, emplacement=emp_a, code='A-07-01')
        self.bin_exp_a = BinLocation.objects.create(
            company=self.a, emplacement=emp_a, code='EXP-07', ordre=990)

        # --- B -----------------------------------------------------------
        self.fb = Fournisseur.objects.create(
            company=self.b, nom='FOURNISSEUR-B-SECRET')
        self.pb = Produit.objects.create(
            company=self.b, nom='ProdB-secret', sku='ASTK7-PB',
            prix_achat=Decimal('10'), prix_vente=Decimal('20'),
            quantite_stock=5)
        self.client_b = Client.objects.create(
            company=self.b, nom='CLIENT-B-SECRET')
        self.cat_b = Categorie.objects.create(company=self.b, nom='Cat B')
        emp_b = EmplacementStock.objects.create(
            company=self.b, nom='Dépôt B', is_principal=True)
        self.bin_b = BinLocation.objects.create(
            company=self.b, emplacement=emp_b, code='B-07-SECRET')

    def assertObjetInexistant(self, reponse, champ):
        self.assertEqual(reponse.status_code, 400, reponse.content)
        self.assertIn(champ, reponse.data, reponse.content)
        self.assertEqual(reponse.data[champ][0].code, 'does_not_exist',
                         reponse.content)

    # --- incident qualité fournisseur (FOUR-1) ------------------------------
    def test_incident_fournisseur_etranger(self):
        corps = {'type_incident': 'autre', 'gravite': 'mineure',
                 'date_incident': JOUR}
        r = self.api.post(f'{BASE}/incidents-qualite-fournisseur/',
                          {**corps, 'fournisseur': self.fb.id}, format='json')
        self.assertObjetInexistant(r, 'fournisseur')
        self.assertNotIn('FOURNISSEUR-B-SECRET', r.content.decode('utf-8'))
        inconnu = self.api.post(f'{BASE}/incidents-qualite-fournisseur/',
                                {**corps, 'fournisseur': 99999999},
                                format='json')
        self.assertObjetInexistant(inconnu, 'fournisseur')
        self.assertFalse(IncidentQualiteFournisseur.objects.exists())
        # Le fournisseur de B reste supprimable (aucun incident de A ne le
        # PROTÈGE).
        self.fb.delete()
        self.assertFalse(Fournisseur.objects.filter(pk=self.fb.pk).exists())

    # --- accord RFA (TEN-6(8)) ----------------------------------------------
    def test_rfa_fournisseur_etranger(self):
        r = self.api.post(f'{BASE}/accords-rfa-fournisseur/', {
            'fournisseur': self.fb.id, 'periode_debut': '2026-01-01',
            'periode_fin': '2026-12-31', 'taux_pct': '2'}, format='json')
        self.assertObjetInexistant(r, 'fournisseur')

    # --- consignation (TEN-6(5)) --------------------------------------------
    def test_consignation_client_etranger(self):
        depot = creer_depot_consignation(
            company=self.a, user=self.resp_a, client_id=self.client_a.id,
            produit_id=self.pa.id, quantite=2, date_depot=JOUR)
        r = self.api.patch(f'{BASE}/consignations/{depot.id}/', {
            'client': self.client_b.id}, format='json')
        self.assertObjetInexistant(r, 'client')
        depot.refresh_from_db()
        self.assertEqual(depot.client_id, self.client_a.id)

        avant = DepotConsignation.objects.count()
        with self.assertRaises(ValueError):
            creer_depot_consignation(
                company=self.a, user=self.resp_a, client_id=self.client_b.id,
                produit_id=self.pa.id, quantite=1, date_depot=JOUR)
        r = self.api.post(f'{BASE}/consignations/', {
            'client': self.client_b.id, 'produit': self.pa.id,
            'quantite_deposee': 1, 'date_depot': JOUR}, format='json')
        self.assertEqual(r.status_code, 400, r.content)
        self.assertEqual(DepotConsignation.objects.count(), avant)
        self.pa.refresh_from_db()
        self.assertEqual(self.pa.quantite_stock, 48)

    def test_export_consignations_jamais_un_client_de_b(self):
        from openpyxl import load_workbook

        depot = creer_depot_consignation(
            company=self.a, user=self.resp_a, client_id=self.client_a.id,
            produit_id=self.pa.id, quantite=2, date_depot=JOUR)
        # État hérité d'avant la borne : client de B injecté en base.
        DepotConsignation.objects.filter(pk=depot.pk).update(
            client=self.client_b)
        r = self.api.get(f'{BASE}/consignations/export-xlsx/')
        self.assertEqual(r.status_code, 200)
        feuille = load_workbook(io.BytesIO(r.content)).active
        valeurs = [str(c.value) for ligne in feuille.iter_rows()
                   for c in ligne if c.value is not None]
        self.assertNotIn('CLIENT-B-SECRET', valeurs)

    # --- autres sérialiseurs déclarés dans les vues -------------------------
    def test_hazmat_casier_etranger(self):
        r = self.api.post(f'{BASE}/casiers-hazmat/', {
            'bin': self.bin_b.id, 'classe_danger': 'BATTERIE_LITHIUM'},
            format='json')
        self.assertObjetInexistant(r, 'bin')

    def test_onboarding_fournisseur_etranger(self):
        r = self.api.post(f'{BASE}/dossiers-onboarding-fournisseur/', {
            'fournisseur': self.fb.id}, format='json')
        self.assertObjetInexistant(r, 'fournisseur')

    def test_profil_saisonnier_produit_etranger(self):
        r = self.api.post(f'{BASE}/profils-saisonniers/', {
            'produit': self.pb.id, 'mois_debut': 1, 'mois_fin': 3,
            'seuil_min': 1, 'nom': 'Hiver'}, format='json')
        self.assertObjetInexistant(r, 'produit')

    def test_plan_echantillonnage_categorie_etrangere(self):
        r = self.api.post(f'{BASE}/plans-echantillonnage/', {
            'categorie': self.cat_b.id, 'taux_echantillon_pct': 10},
            format='json')
        self.assertObjetInexistant(r, 'categorie')

    def test_reappro_casier_etranger(self):
        r = self.api.post(f'{BASE}/seuils-reappro-casier/', {
            'bin': self.bin_b.id, 'produit': self.pa.id, 'seuil': 1,
            'quantite_cible': 5}, format='json')
        self.assertObjetInexistant(r, 'bin')
        r = self.api.post(f'{BASE}/taches-reappro-interne/', {
            'produit': self.pb.id, 'bin_cible': self.bin_a.id,
            'quantite': 1}, format='json')
        self.assertObjetInexistant(r, 'produit')

    # --- retour fournisseur scanné : bins_source bornés ---------------------
    def test_valider_scanne_casier_source_etranger(self):
        retour = RetourFournisseur.objects.create(
            company=self.a, reference='RET-ASTK7-0001', fournisseur=self.fa)
        ligne = LigneRetourFournisseur.objects.create(
            retour=retour, produit=self.pa, quantite=1, motif='Défaut')
        r = self.api.post(
            f'{BASE}/retours-fournisseur/{retour.id}/valider-scanne/',
            {'bins_source': {str(ligne.id): self.bin_b.id}}, format='json')
        self.assertEqual(r.status_code, 400, r.content)
        retour.refresh_from_db()
        self.assertEqual(retour.statut, RetourFournisseur.Statut.BROUILLON)
        self.assertFalse(MouvementStock.objects.filter(
            reference='RET-ASTK7-0001').exists())
        self.pa.refresh_from_db()
        self.assertEqual(self.pa.quantite_stock, 50)

    # --- ENF17 : FK en lecture seule des sérialiseurs restants ---------------
    def test_enf17_fk_lecture_seule_jamais_ecrites(self):
        """Un id de B posté sur ces FK en lecture seule n'est jamais écrit ;
        chaque sérialiseur porte le mixin qui les bornerait sinon."""
        from types import SimpleNamespace

        from core.serializers import CompanyScopedRelationsMixin
        from apps.stock.serializers import (
            LotEntrepotSerializer, RevisionKitSerializer,
        )
        from apps.stock.views.budget_departement import (
            EngagementBudgetSerializer,
        )
        from apps.stock.views.catalogue_achat import CatalogueAchatSerializer

        ctx = {'request': SimpleNamespace(user=self.resp_a)}
        cas = (
            (RevisionKitSerializer, {'kit': self.pb.pk, 'user': 1}),
            (LotEntrepotSerializer, {'produit': self.pb.pk,
                                     'emplacement': self.bin_b.emplacement_id}),
            (EngagementBudgetSerializer, {'budget': 1, 'demande_achat': 1,
                                          'bon_commande': 1}),
            (CatalogueAchatSerializer, {'categorie': self.cat_b.pk}),
        )
        for cls, corps in cas:
            with self.subTest(serializer=cls.__name__):
                self.assertTrue(issubclass(cls, CompanyScopedRelationsMixin))
                ser = cls(data=corps, partial=True, context=ctx)
                self.assertTrue(ser.is_valid(), ser.errors)
                for champ in corps:
                    self.assertNotIn(champ, ser.validated_data)
