"""ASTK13 (D-ASTK-2) — montants d'achat de la fiche fournisseur (vue-360,
performance, export-conformite JSON et xlsx) servis UNIQUEMENT avec
`prix_achat_voir`.

Rejoue FOUR-3 / TEN-8(b) : pour un rôle ['stock_voir'], vue-360 = 200 avec
accords_prix=[{prix_convenu:'780.00'}] et solde_total_du=5000.00,
performance = 200 avec total_achats_ht=10000.00, export-conformite = 200 avec
montant_achete=[10000.00].

Source réelle : selectors/services réels (aucun mock).

Run :
    python manage.py test apps.stock.test_astk_prix_achat_fournisseur -v 2
"""
import datetime
import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import AccessToken

from apps.roles.models import Role
from apps.roles.permissions_registre import CANONICAL_SYSTEM_ROLES
from apps.stock.models import (
    BonCommandeFournisseur, FactureFournisseur, Fournisseur,
    LigneBonCommandeFournisseur, PalierPrixFournisseur, PrixFournisseur,
    Produit,
)
from authentication.models import Company

User = get_user_model()

BASE = '/api/django/stock/fournisseurs/'


def _api(user):
    api = APIClient()
    api.credentials(HTTP_AUTHORIZATION=f'Bearer {AccessToken.for_user(user)}')
    return api


class PrixAchatFournisseurTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(nom='ASTK13', slug='astk13-co')
        role = Role.objects.create(
            company=self.company, nom='r-astk13-lecture',
            permissions=['stock_voir'])
        self.lecteur = User.objects.create_user(
            username='astk13-lecteur', password='x', company=self.company,
            role=role)
        self.assertFalse(self.lecteur.can_view_buy_prices)
        role_admin = Role.objects.create(
            company=self.company, nom='Administrateur',
            permissions=dict(CANONICAL_SYSTEM_ROLES)['Administrateur'])
        self.admin = User.objects.create_user(
            username='astk13-admin', password='x', company=self.company,
            role=role_admin)
        self.assertTrue(self.admin.can_view_buy_prices)

        self.fournisseur = Fournisseur.objects.create(
            company=self.company, nom='Fournisseur ASTK13')
        self.produit = Produit.objects.create(
            company=self.company, nom='Onduleur ASTK13', sku='ASTK13-1',
            prix_vente=Decimal('1500'), prix_achat=Decimal('900'))
        prix = PrixFournisseur.objects.create(
            company=self.company, produit=self.produit,
            fournisseur=self.fournisseur, prix_achat=Decimal('820.00'))
        PalierPrixFournisseur.objects.create(
            prix_fournisseur=prix, qte_min=10, prix=Decimal('780.00'))
        bc = BonCommandeFournisseur.objects.create(
            company=self.company, reference='BCF-ASTK13-1',
            fournisseur=self.fournisseur,
            statut=BonCommandeFournisseur.Statut.ENVOYE,
            date_commande=datetime.date.today())
        LigneBonCommandeFournisseur.objects.create(
            bon_commande=bc, produit=self.produit, quantite=10,
            prix_achat_unitaire=Decimal('1000'))
        FactureFournisseur.objects.create(
            company=self.company, reference='FF-ASTK13-1',
            fournisseur=self.fournisseur,
            date_facture=datetime.date.today(),
            montant_ttc=Decimal('5000'))

    def _conformite_xlsx(self, user):
        rep = _api(user).get(f'{BASE}export-conformite/', {'export': 'xlsx'})
        self.assertEqual(rep.status_code, 200)
        ws = load_workbook(io.BytesIO(rep.content)).active
        return [list(r) for r in ws.iter_rows(values_only=True)]

    def test_vue360_sans_prix(self):
        rep = _api(self.lecteur).get(f'{BASE}{self.fournisseur.pk}/vue-360/')
        self.assertEqual(rep.status_code, 200)
        self.assertNotIn('solde_total_du', rep.data)
        self.assertEqual(len(rep.data['accords_prix']), 1)
        self.assertNotIn('prix_convenu', rep.data['accords_prix'][0])
        # Le reste de la fiche est intact.
        self.assertEqual(rep.data['factures_ouvertes'], 1)
        self.assertEqual(rep.data['bcf_ouverts'], 1)
        self.assertEqual(rep.data['accords_prix_actifs'], 1)

    def test_performance_sans_prix(self):
        rep = _api(self.lecteur).get(
            f'{BASE}{self.fournisseur.pk}/performance/')
        self.assertEqual(rep.status_code, 200)
        self.assertNotIn('total_achats_ht', rep.data)
        self.assertEqual(rep.data['nb_bons'], 1)

    def test_export_conformite_sans_montant(self):
        rep = _api(self.lecteur).get(f'{BASE}export-conformite/')
        self.assertEqual(rep.status_code, 200)
        self.assertEqual(len(rep.data), 1)
        self.assertNotIn('montant_achete', rep.data[0])
        lignes = self._conformite_xlsx(self.lecteur)
        self.assertNotIn('Montant acheté (MAD HT)', lignes[0])
        self.assertNotIn(10000, [c for row in lignes for c in row])

    def test_administrateur_inchange(self):
        api = _api(self.admin)
        rep = api.get(f'{BASE}{self.fournisseur.pk}/vue-360/')
        self.assertEqual(Decimal(rep.data['solde_total_du']),
                         Decimal('5000'))
        self.assertEqual(rep.data['accords_prix'][0]['prix_convenu'],
                         '780.00')
        rep = api.get(f'{BASE}{self.fournisseur.pk}/performance/')
        self.assertEqual(Decimal(rep.data['total_achats_ht']),
                         Decimal('10000'))
        rep = api.get(f'{BASE}export-conformite/')
        self.assertEqual(Decimal(str(rep.data[0]['montant_achete'])),
                         Decimal('10000'))
        lignes = self._conformite_xlsx(self.admin)
        self.assertIn('Montant acheté (MAD HT)', lignes[0])
