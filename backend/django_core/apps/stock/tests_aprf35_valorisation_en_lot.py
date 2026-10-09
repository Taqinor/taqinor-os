"""APRF35 — la valorisation de TOUT le catalogue se calcule en lot : le nombre
de requêtes de ``stock_valuation_by_location`` (écran Valorisation et export
xlsx) ne dépend plus du nombre de produits, et chaque coût calculé en lot est
égal AU CENTIME (coût, source) à l'accesseur unitaire
``valuation_cost_with_source`` — méthodes coût moyen ET FIFO. L'export des
consignations (négoce) lit la dernière déclaration de chaque dépôt sans une
requête par dépôt.

Données : revalorisations validées (avec réceptions avant/après), lignes de
BCF reçues avec frais annexes, entrées de production à ``cout_unitaire``,
stock ventilé sur un second emplacement et marchandise en transit."""
import datetime
import io
from decimal import Decimal
from unittest import mock

from django.contrib.auth import get_user_model
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import AccessToken

from apps.crm.models import Client
from apps.stock.models import (
    BonCommandeFournisseur, EmplacementStock, Fournisseur,
    LigneBonCommandeFournisseur, MouvementStock, Produit,
    RevalorisationStock, StockEmplacement, TransfertStock,
)
from apps.stock.models_consignation import (
    DeclarationConsommation, DepotConsignation,
)
from apps.stock.services import (
    VALUATION_FIFO, VALUATION_WAVG, ensure_emplacements,
    stock_valuation_by_location, valuation_cost_with_source,
    valuation_costs_for,
)
from authentication.models import Company

User = get_user_model()
VALO = '/api/django/stock/produits/valorisation/'
VALO_XLSX = '/api/django/stock/produits/valorisation-xlsx/'
CONSIGNATIONS_XLSX = '/api/django/stock/consignations/export-xlsx/'


def _api(user):
    api = APIClient()
    api.credentials(HTTP_AUTHORIZATION=f'Bearer {AccessToken.for_user(user)}')
    return api


class ValorisationEnLotTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(nom='APRF35', slug='aprf35-co')
        self.user = User.objects.create_user(
            username='aprf35', password='x', company=self.company,
            role_legacy='admin')
        self.fournisseur = Fournisseur.objects.create(
            company=self.company, nom='Four APRF35')
        ensure_emplacements(self.company)
        self.principal = EmplacementStock.objects.get(
            company=self.company, is_principal=True)
        self.camion = EmplacementStock.objects.filter(
            company=self.company, is_principal=False).first()
        self.aujourdhui = timezone.localdate()
        self.n = 0

    # ── fabrique de données ─────────────────────────────────────────────
    def _recu(self, produit, quantite, prix, il_y_a_jours, frais='0'):
        self.n += 1
        bc = BonCommandeFournisseur.objects.create(
            company=self.company, reference=f'BCF-APRF35-{self.n}',
            fournisseur=self.fournisseur,
            statut=BonCommandeFournisseur.Statut.RECU)
        jour = self.aujourdhui - datetime.timedelta(days=il_y_a_jours)
        BonCommandeFournisseur.objects.filter(pk=bc.pk).update(
            date_creation=timezone.make_aware(
                datetime.datetime.combine(jour, datetime.time(10, 0))))
        LigneBonCommandeFournisseur.objects.create(
            bon_commande=bc, produit=produit, quantite=quantite,
            prix_achat_unitaire=Decimal(prix), quantite_recue=quantite,
            frais_annexes=Decimal(frais))

    def _produits(self, nombre):
        crees = []
        for _ in range(nombre):
            self.n += 1
            i = self.n
            p = Produit.objects.create(
                company=self.company, nom=f'Produit APRF35 {i:04d}',
                sku=f'APRF35-{i}', prix_vente=Decimal('500'),
                prix_achat=Decimal('111.11'), quantite_stock=12)
            genre = i % 4
            if genre == 0:
                self._recu(p, 10, '100', 20, frais='30')
                self._recu(p, 5, '130.33', 3)
            elif genre == 1:
                self._recu(p, 10, '1000', 30)
                RevalorisationStock.objects.create(
                    company=self.company, produit=p,
                    ancien_cout=Decimal('1000'), nouveau_cout=Decimal('600'),
                    quantite_snapshot=10, delta_valeur=Decimal('-4000'),
                    motif='Baisse', auteur=self.user,
                    statut=RevalorisationStock.Statut.VALIDEE,
                    date_validation=timezone.now() - datetime.timedelta(
                        days=10))
                self._recu(p, 4, '700', 2)
            elif genre == 2:
                mvt = MouvementStock.objects.create(
                    company=self.company, produit=p, type_mouvement='entree',
                    quantite=6, quantite_avant=6, quantite_apres=12,
                    cout_unitaire=Decimal('87.1234'))
                MouvementStock.objects.filter(pk=mvt.pk).update(
                    date=timezone.now() - datetime.timedelta(days=1))
                self._recu(p, 3, '90', 4)
            # genre 3 : aucun achat → prix catalogue.
            if i % 3 == 0:
                StockEmplacement.objects.create(
                    company=self.company, produit=p, emplacement=self.camion,
                    quantite=4)
            if i % 5 == 0:
                TransfertStock.objects.create(
                    company=self.company, produit=p, source=self.principal,
                    destination=self.camion, quantite=2,
                    statut=TransfertStock.Statut.EXPEDIE)
            crees.append(p)
        return crees

    def _compter(self, fn):
        with CaptureQueriesContext(connection) as ctx:
            fn()
        return len(ctx.captured_queries)

    # ── taille : requêtes constantes ────────────────────────────────────
    def test_requetes_constantes_cout_moyen(self):
        self._produits(9)
        petit = self._compter(
            lambda: stock_valuation_by_location(self.company))
        self._produits(20)
        grand = self._compter(
            lambda: stock_valuation_by_location(self.company))
        self.assertEqual(petit, grand)

    def test_requetes_constantes_fifo(self):
        with mock.patch('apps.stock.services.stock_valuation_method',
                        return_value=VALUATION_FIFO):
            self._produits(9)
            petit = self._compter(
                lambda: stock_valuation_by_location(self.company))
            self._produits(20)
            grand = self._compter(
                lambda: stock_valuation_by_location(self.company))
        self.assertEqual(petit, grand)

    def test_ecran_et_export_requetes_constantes(self):
        self._produits(9)
        api = _api(self.user)
        ecran_petit = self._compter(lambda: api.get(VALO))
        export_petit = self._compter(lambda: api.get(VALO_XLSX))
        self._produits(20)
        ecran_grand = self._compter(lambda: api.get(VALO))
        export_grand = self._compter(lambda: api.get(VALO_XLSX))
        self.assertEqual(ecran_petit, ecran_grand)
        self.assertEqual(export_petit, export_grand)
        self.assertEqual(api.get(VALO).status_code, 200)

    # ── croisé : lot == unitaire, au centime ────────────────────────────
    def _croise(self, method):
        produits = self._produits(24)
        lot = valuation_costs_for(
            list(Produit.objects.filter(company=self.company)), method)
        for p in produits:
            p = Produit.objects.get(pk=p.pk)
            self.assertEqual(
                lot[p.id], valuation_cost_with_source(p, method=method),
                f'{p.nom} ({method})')
        # Le lot couvre bien les trois sources de coût.
        sources = {source for _cout, source in lot.values()}
        self.assertTrue({'achats', 'revalorisation', 'catalogue'} <= sources)

    def test_lot_egal_unitaire_cout_moyen(self):
        self._croise(VALUATION_WAVG)

    def test_lot_egal_unitaire_fifo(self):
        self._croise(VALUATION_FIFO)

    def test_lignes_ecran_egales_unitaire_et_transit_exclu_du_principal(self):
        produits = self._produits(12)
        data = stock_valuation_by_location(self.company)
        for p in produits:
            cout, source = valuation_cost_with_source(
                Produit.objects.get(pk=p.pk), method=VALUATION_WAVG)
            lignes = [x for x in data['lignes'] if x['produit_id'] == p.id]
            self.assertTrue(lignes)
            for ligne in lignes:
                self.assertEqual(ligne['cout_moyen'], cout)
                self.assertEqual(ligne['source'], source)
            # 12 en stock ventilés (principal + camion + transit) sans perte.
            self.assertEqual(sum(x['quantite'] for x in lignes), 12)

    # ── export des consignations : dernière déclaration sans N+1 ────────
    def _depots(self, nombre):
        client = Client.objects.create(
            company=self.company, nom=f'Client APRF35 {self.n}', prenom='Test')
        produit = Produit.objects.create(
            company=self.company, nom='Conso APRF35', sku=f'APRF35-C{self.n}',
            prix_vente=Decimal('1'), quantite_stock=100)
        self.n += 1
        attendu = {}
        for k in range(nombre):
            depot = DepotConsignation.objects.create(
                company=self.company, client=client, produit=produit,
                quantite_deposee=10, date_depot=datetime.date(2026, 1, 1),
                adresse_site=f'Site {self.n}-{k}')
            for jour in (3, 9 + k):
                DeclarationConsommation.objects.create(
                    company=self.company, depot=depot, quantite=1,
                    date_declaration=datetime.date(2026, 2, jour))
            attendu[depot.adresse_site] = datetime.date(
                2026, 2, 9 + k).isoformat()
        return attendu

    def test_export_consignations_requetes_constantes(self):
        api = _api(self.user)
        attendu = self._depots(2)
        petit = self._compter(lambda: api.get(CONSIGNATIONS_XLSX))
        attendu.update(self._depots(6))
        grand = self._compter(lambda: api.get(CONSIGNATIONS_XLSX))
        self.assertEqual(petit, grand)
        resp = api.get(CONSIGNATIONS_XLSX)
        self.assertEqual(resp.status_code, 200)
        ws = load_workbook(io.BytesIO(resp.content)).active
        lues = {row[3]: row[7] for row in ws.iter_rows(
            min_row=2, values_only=True) if row[0] != 'TOTAL'}
        self.assertEqual(lues, attendu)
