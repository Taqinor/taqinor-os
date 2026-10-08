"""ASTK208 (C-ASTK-053, volet MVT-23) — le « Net » de l'historique des
mouvements = Σ(quantité après − quantité avant) : il réconcilie avec le
stock ; rebuts et ajustements (signés) ont leurs colonnes.

Sonde MVT-23 d'origine : entrée 10 puis rebut 5 (stock réel 5) → agrégation
(entrées 10, sorties 0, net 10) ; un ajustement −2 ne comptait nulle part.

Aucun mock : sélecteur et vue réels.

Run :
    python manage.py test apps.stock.test_astk_mouvements_net -v 2
"""
import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import AccessToken

from authentication.models import Company
from apps.stock.models import MouvementStock, Produit
from apps.stock.selectors import mouvements_agreges

User = get_user_model()

URL = '/api/django/stock/mouvements/agregation/'


def _api(user):
    api = APIClient()
    api.credentials(HTTP_AUTHORIZATION=f'Bearer {AccessToken.for_user(user)}')
    return api


class NetTests(TestCase):
    def setUp(self):
        self.co = Company.objects.create(nom='astk208', slug='astk208')
        self.admin = User.objects.create_user(
            username='astk208-admin', password='x', company=self.co,
            role_legacy='admin')
        self.pa = Produit.objects.create(
            company=self.co, nom='Pa ASTK208', sku='ASTK208-A',
            prix_vente=Decimal('20'), prix_achat=Decimal('10'),
            quantite_stock=5)
        self.pb = Produit.objects.create(
            company=self.co, nom='Pb ASTK208', sku='ASTK208-B',
            prix_vente=Decimal('20'), prix_achat=Decimal('10'),
            quantite_stock=8)
        T = MouvementStock.TypeMouvement
        self._mvt(self.pa, T.ENTREE, 10, 0, 10)
        self._mvt(self.pa, T.REBUT, 5, 10, 5)
        self._mvt(self.pb, T.AJUSTEMENT, 2, 10, 8)

    def _mvt(self, produit, type_mouvement, quantite, avant, apres):
        return MouvementStock.objects.create(
            company=self.co, produit=produit, type_mouvement=type_mouvement,
            quantite=quantite, quantite_avant=avant, quantite_apres=apres)

    def _par_libelle(self, rows):
        return {r['libelle']: r for r in rows}

    def test_rebut_compte_dans_le_net(self):
        rows = self._par_libelle(
            mouvements_agreges(self.co, group_by='produit'))
        self.assertEqual(rows['Pa ASTK208'], {
            'libelle': 'Pa ASTK208', 'entrees': 10, 'sorties': 0,
            'rebuts': 5, 'ajustements': 0, 'net': 5})
        self.pa.refresh_from_db()
        self.assertEqual(rows['Pa ASTK208']['net'], self.pa.quantite_stock)

    def test_ajustement_signe(self):
        rows = self._par_libelle(
            mouvements_agreges(self.co, group_by='produit'))
        self.assertEqual(rows['Pb ASTK208']['ajustements'], -2)
        self.assertEqual(rows['Pb ASTK208']['net'], -2)

    def test_endpoint_et_export(self):
        api = _api(self.admin)
        rep = api.get(URL, {'group_by': 'produit'})
        self.assertEqual(rep.status_code, 200, rep.content)
        pa = next(r for r in rep.json() if r['libelle'] == 'Pa ASTK208')
        self.assertEqual((pa['rebuts'], pa['ajustements'], pa['net']),
                         (5, 0, 5))
        rep = api.get(URL, {'group_by': 'produit', 'export': 'xlsx'})
        self.assertEqual(rep.status_code, 200)
        feuille = load_workbook(io.BytesIO(rep.content)).active
        lignes = list(feuille.iter_rows(values_only=True))
        self.assertEqual(
            list(lignes[0]),
            ['Groupe', 'Entrées', 'Sorties', 'Rebuts', 'Ajustements', 'Net'])
        self.assertIn(('Pa ASTK208', 10, 0, 5, 0, 5),
                      [tuple(ligne) for ligne in lignes[1:]])
