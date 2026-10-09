"""APRF34 — l'export xlsx des mouvements bascule en job de fond au-delà du
seuil NTPLT30 (202 + id de job, livrable MinIO, notification « export prêt »)
et reste synchrone — fichier inchangé — en dessous.

Le seuil est abaissé par ``NTPLT30_EXPORT_ROW_THRESHOLD`` ; la tâche Celery
est exécutée en mode eager (``delay`` → appel direct) ; MinIO est remplacé
par un dépôt en mémoire (aucun conteneur)."""
import io
import os
from datetime import timedelta
from decimal import Decimal
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import AccessToken

from apps.notifications.models import Notification
from apps.stock.models import MouvementStock, Produit
from apps.stock.tasks import export_mouvements_xlsx_task
from authentication.models import Company
from core.models import BackgroundJob

User = get_user_model()
URL = '/api/django/stock/mouvements/export-xlsx/'


def _api(user):
    api = APIClient()
    api.credentials(HTTP_AUTHORIZATION=f'Bearer {AccessToken.for_user(user)}')
    return api


def _cells(data):
    ws = load_workbook(io.BytesIO(data)).active
    return [[c.value for c in row] for row in ws.iter_rows()]


class ExportMouvementsAsyncTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(nom='APRF34', slug='aprf34-co')
        self.user = User.objects.create_user(
            username='aprf34', password='x', company=self.company,
            role_legacy='admin')
        self.produit = Produit.objects.create(
            company=self.company, nom='Panneau APRF34', sku='APRF34-1',
            prix_vente=Decimal('100'), prix_achat=Decimal('70'),
            quantite_stock=0)
        autre = Company.objects.create(nom='Autre', slug='aprf34-autre')
        p_autre = Produit.objects.create(
            company=autre, nom='Etranger', sku='APRF34-X',
            prix_vente=Decimal('1'), quantite_stock=0)
        MouvementStock.objects.create(
            company=autre, produit=p_autre, type_mouvement='entree',
            quantite=1, quantite_avant=0, quantite_apres=1)
        self.stockes = {}

    def _mouvements(self, n):
        MouvementStock.objects.bulk_create([
            MouvementStock(
                company=self.company, produit=self.produit,
                type_mouvement='entree' if i % 2 else 'sortie',
                quantite=i + 1, quantite_avant=i, quantite_apres=2 * i + 1,
                reference=f'MV-{i}', note='=SOMME(A1)' if i == 0 else 'n',
                created_by=self.user)
            for i in range(n)])
        # Dates distinctes : l'ordre « -date » de la liste est alors total.
        base = timezone.now()
        for i, pk in enumerate(MouvementStock.objects.filter(
                company=self.company).order_by('pk').values_list(
                    'pk', flat=True)):
            MouvementStock.objects.filter(pk=pk).update(
                date=base - timedelta(minutes=i))

    def _store(self, data, *, company_id, job_id, ext='xlsx',
               content_type=None):
        key = f'exports/{company_id}/{job_id}.{ext}'
        self.stockes[key] = data
        return key

    def _export(self, seuil):
        eager = mock.patch.object(
            export_mouvements_xlsx_task, 'delay',
            side_effect=lambda **kw: export_mouvements_xlsx_task(**kw))
        with mock.patch.dict(os.environ,
                             {'NTPLT30_EXPORT_ROW_THRESHOLD': str(seuil)}), \
                eager, \
                mock.patch('apps.records.storage.store_export_result',
                           side_effect=self._store), \
                mock.patch('apps.records.storage.presign_export_result',
                           return_value='https://minio.test/signe'):
            return _api(self.user).post(URL)

    def test_sous_le_seuil_reste_synchrone(self):
        self._mouvements(5)
        resp = self._export(10)
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(BackgroundJob.objects.filter(
            company=self.company).exists())
        lignes = _cells(resp.content)
        self.assertEqual(lignes[0][:3], ['Référence', 'Type', 'Produit'])
        self.assertEqual(len(lignes), 6)  # en-tête + 5, pas l'autre société

    def test_au_dela_du_seuil_202_et_job(self):
        self._mouvements(12)
        resp = self._export(10)
        self.assertEqual(resp.status_code, 202)
        corps = resp.json()
        self.assertIn('job_id', corps)
        self.assertIn('statut', corps)
        job = BackgroundJob.objects.get(pk=corps['job_id'])
        self.assertEqual(job.company_id, self.company.id)
        self.assertEqual(job.user_id, self.user.id)
        self.assertEqual(job.statut, BackgroundJob.STATUT_DONE)
        self.assertEqual(job.result_file_key,
                         f'exports/{self.company.id}/{job.pk}.xlsx')
        notif = Notification.objects.filter(recipient=self.user).first()
        self.assertIsNotNone(notif)
        self.assertIn('https://minio.test/signe', notif.body)
        self.assertEqual(notif.link, 'https://minio.test/signe')

    def test_meme_contenu_que_l_export_synchrone(self):
        self._mouvements(12)
        synchrone = self._export(0)  # 0 = jamais asynchrone (off-switch)
        self.assertEqual(synchrone.status_code, 200)
        avant = MouvementStock.objects.filter(
            company=self.company).values_list('pk', 'quantite_apres')
        avant = sorted(avant)
        resp = self._export(10)
        self.assertEqual(resp.status_code, 202)
        job = BackgroundJob.objects.get(pk=resp.json()['job_id'])
        livrable = self.stockes[job.result_file_key]
        self.assertEqual(_cells(livrable), _cells(synchrone.content))
        # Lecture seule : aucun mouvement modifié.
        self.assertEqual(sorted(MouvementStock.objects.filter(
            company=self.company).values_list('pk', 'quantite_apres')), avant)
