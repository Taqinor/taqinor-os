"""AANA12 (C-AANA-031) — chaque ligne d'import produit est VALIDÉE.

Rouge figé par l'audit (05/10) sur le CSV ``X,S1,abc`` / ``Y,S2,-5`` /
``Z,S3,10,2.7`` : ``skipped=[]``, P1.prix_vente=0 (prix « abc » avalé puis
remplacé par 0), P2.prix_vente=-5, P3.quantite=2 (2.7 tronqué). Et une cellule
xlsx numérique entière stockée ``612345678.0`` arrivait en téléphone
``'612345678.0'``. Correctif : erreur par ligne dans ``skipped``, aucune
valeur coercée en silence, cellule entière rendue sans « .0 »."""
import io
import zipfile

from django.core.files.uploadedfile import SimpleUploadedFile

from apps.crm.models import Lead
from apps.stock.models import Produit

from .parsing import iter_rows
from .tests import ImportBase


def _xlsx_avec_flottant_entier():
    """Classeur RÉEL dont la cellule téléphone est sérialisée ``612345678.0``
    (forme écrite par plusieurs tableurs ; openpyxl écrit ``612345678``)."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(['Nom', 'Telephone'])
    ws.append(['Xlsx Lead', 612345678])
    tampon = io.BytesIO()
    wb.save(tampon)
    entree = zipfile.ZipFile(io.BytesIO(tampon.getvalue()))
    sortie_buf = io.BytesIO()
    with zipfile.ZipFile(sortie_buf, 'w', zipfile.ZIP_DEFLATED) as sortie:
        for item in entree.infolist():
            data = entree.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                assert b'<v>612345678</v>' in data
                data = data.replace(b'<v>612345678</v>',
                                    b'<v>612345678.0</v>')
            sortie.writestr(item, data)
    return sortie_buf.getvalue()


class TestValidationLignesProduits(ImportBase):
    def _importer(self, contenu):
        resp = self.api_dir.post('/api/django/imports/commit/', {
            'file': self._csv(contenu), 'target': 'products',
        }, format='multipart')
        self.assertEqual(resp.status_code, 200, resp.data)
        return resp.data

    def test_prix_invalide_signale(self):
        data = self._importer('Nom,SKU,Prix,Quantite\n'
                              'X,S1,abc,\n'
                              'Y,S2,-5,\n'
                              'Z,S3,10,2.7\n'
                              'W,S4,"1 200,50",3\n')
        lignes = {s['ligne']: s['raison'] for s in data['skipped']}
        self.assertEqual(sorted(lignes), [1, 2, 3], data)
        self.assertIn('abc', lignes[1])
        self.assertIn('négatif', lignes[2])
        self.assertIn('non entière', lignes[3])
        self.assertEqual(data['created'], 1, data)
        self.assertFalse(Produit.objects.filter(
            company=self.company, sku__in=['S1', 'S2', 'S3']).exists())
        w = Produit.objects.get(company=self.company, sku='S4')
        w.refresh_from_db()
        self.assertEqual(str(w.prix_vente), '1200.50')
        self.assertEqual(w.quantite_stock, 3)

    def test_quantite_entiere_ecrite_en_decimal_acceptee(self):
        data = self._importer('Nom,SKU,Prix,Quantite\nV,S5,10,4.0\n')
        self.assertEqual(data['created'], 1, data)
        self.assertEqual(
            Produit.objects.get(company=self.company, sku='S5').quantite_stock,
            4)

    def test_apercu_n_annonce_pas_la_creation_refusee(self):
        resp = self.api_dir.post('/api/django/imports/dry-run/', {
            'file': self._csv('Nom,SKU,Prix\nX,S1,abc\nW,S4,10\n'),
            'target': 'products',
        }, format='multipart')
        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(resp.data['resume']['creation'], 1, resp.data)
        self.assertEqual(resp.data['resume']['ignoree'], 1, resp.data)


class TestXlsxEntierSansPointZero(ImportBase):
    def test_parseur_rend_l_entier(self):
        _, rows = iter_rows(_xlsx_avec_flottant_entier(), 'f.xlsx')
        self.assertEqual(rows[0]['Telephone'], 612345678)

    def test_telephone_xlsx_sans_point_zero(self):
        fichier = SimpleUploadedFile(
            'leads.xlsx', _xlsx_avec_flottant_entier(),
            content_type='application/vnd.openxmlformats-officedocument.'
                         'spreadsheetml.sheet')
        resp = self.api.post('/api/django/imports/commit/', {
            'file': fichier, 'target': 'leads',
        }, format='multipart')
        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(resp.data['created'], 1, resp.data)
        lead = Lead.objects.get(company=self.company, nom='Xlsx Lead')
        self.assertEqual(lead.telephone, '612345678')
