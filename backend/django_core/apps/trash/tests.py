"""NTUX7 — Tests de la corbeille transverse 30 jours (`apps.trash`).

Couvre : journalisation par l'ÉVÉNEMENT `record_soft_deleted` (l'émetteur ne
connaît pas la corbeille), société posée côté serveur + isolation multi-société,
garde Directeur/Admin de l'écran `/parametres/corbeille`, restauration via le
restaurateur enregistré par l'app cible PUIS via le repli générique, journal en
lecture seule, et purge de rétention.

Le lead CRM sert de cible RÉELLE (critère d'acceptation NTUX7 : « archiver un
lead puis le restaurer depuis la corbeille redonne l'état exact ») — arête
test→domaine assumée, le code de PRODUCTION de `apps.trash` n'importe aucune
app métier (la cible est résolue par `contenttypes`).
"""
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.contrib.contenttypes.models import ContentType
from django.core.management import call_command
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import AccessToken

from apps.crm.models import Lead
from apps.roles.models import Role
from authentication.models import Company
from core.events import record_soft_deleted

from . import registry
from .models import RETENTION_JOURS, ElementSupprime
from .services import RestaurationImpossible, purger_expires, restaurer

User = get_user_model()


def make_company(slug, nom):
    company, _ = Company.objects.get_or_create(slug=slug, defaults={'nom': nom})
    return company


def make_user(company, username, role_legacy='normal'):
    return User.objects.create_user(
        username=username, password='x', company=company, role_legacy=role_legacy)


def auth(user):
    api = APIClient()
    api.credentials(HTTP_AUTHORIZATION=f'Bearer {AccessToken.for_user(user)}')
    return api


def rows(resp):
    data = resp.data
    return data['results'] if isinstance(data, dict) and 'results' in data else data


def archiver(lead, user=None, **extra):
    """Simule ce que fera une app émettrice : soft-delete + émission."""
    lead.is_archived = True
    lead.save(update_fields=['is_archived'])
    record_soft_deleted.send(
        sender=type(lead), instance=lead, company=lead.company, user=user, **extra)
    return lead


class CorbeilleBase(TestCase):
    BASE = '/api/django/trash/corbeille/'

    def setUp(self):
        self.co_a = make_company('trash-a', 'Trash A')
        self.co_b = make_company('trash-b', 'Trash B')
        self.directeur = make_user(self.co_a, 'trash-directeur', role_legacy='responsable')
        self.commercial = make_user(self.co_a, 'trash-com', role_legacy='normal')
        self.directeur_b = make_user(self.co_b, 'trash-b-directeur', role_legacy='responsable')
        self.lead = Lead.objects.create(company=self.co_a, nom='Alaoui')

    def tearDown(self):
        # Le registre est un état de PROCESSUS : un test qui enregistre un
        # restaurateur ne doit pas contaminer les suivants.
        registry._RESTAURATEURS.clear()


class JournalisationTests(CorbeilleBase):
    def test_evenement_cree_une_entree_de_corbeille(self):
        archiver(self.lead, user=self.directeur,
                 type_libelle='Lead', libelle='Alaoui', donnees={'ville': 'Rabat'})
        element = ElementSupprime.objects.get()
        self.assertEqual(element.company, self.co_a)
        self.assertEqual(element.object_id, self.lead.pk)
        self.assertEqual(element.cle_modele, 'crm.lead')
        self.assertEqual(element.type_libelle, 'Lead')
        self.assertEqual(element.libelle_snapshot, 'Alaoui')
        self.assertEqual(element.donnees_snapshot, {'ville': 'Rabat'})
        self.assertEqual(element.supprime_par, self.directeur)
        self.assertIsNone(element.restaure_le)

    def test_expiration_a_trente_jours(self):
        archiver(self.lead)
        element = ElementSupprime.objects.get()
        ecart = element.expire_le - element.supprime_le
        self.assertEqual(ecart, timedelta(days=RETENTION_JOURS))

    def test_libelles_par_defaut_depuis_la_cible(self):
        archiver(self.lead)
        element = ElementSupprime.objects.get()
        self.assertEqual(element.libelle_snapshot, str(self.lead))
        self.assertTrue(element.type_libelle)

    def test_reemission_ne_duplique_pas_lentree(self):
        archiver(self.lead)
        archiver(self.lead)
        self.assertEqual(ElementSupprime.objects.count(), 1)

    def test_emission_sans_company_est_ignoree(self):
        record_soft_deleted.send(sender=Lead, instance=None, company=None)
        self.assertEqual(ElementSupprime.objects.count(), 0)


class CorbeilleApiTests(CorbeilleBase):
    def test_liste_reservee_directeur_admin(self):
        archiver(self.lead)
        self.assertEqual(auth(self.commercial).get(self.BASE).status_code, 403)
        self.assertEqual(auth(self.directeur).get(self.BASE).status_code, 200)

    def test_isolation_multi_societe(self):
        archiver(self.lead)
        resp = auth(self.directeur_b).get(self.BASE)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(rows(resp)), 0)

    def test_filtre_par_type_et_par_date(self):
        archiver(self.lead, type_libelle='Lead')
        autre = Lead.objects.create(company=self.co_a, nom='Bennani')
        archiver(autre, type_libelle='Devis')
        api = auth(self.directeur)
        self.assertEqual(len(rows(api.get(self.BASE, {'type': 'Lead'}))), 1)
        demain = (timezone.now() + timedelta(days=1)).isoformat()
        self.assertEqual(len(rows(api.get(self.BASE, {'depuis': demain}))), 0)

    def test_journal_en_lecture_seule(self):
        archiver(self.lead)
        element = ElementSupprime.objects.get()
        api = auth(self.directeur)
        self.assertEqual(api.post(self.BASE, {}, format='json').status_code, 405)
        self.assertEqual(api.delete(f'{self.BASE}{element.pk}/').status_code, 405)
        self.assertEqual(
            api.patch(f'{self.BASE}{element.pk}/', {'type_libelle': 'X'},
                      format='json').status_code, 405)

    def test_restaurer_remet_la_cible_active(self):
        archiver(self.lead, user=self.directeur)
        element = ElementSupprime.objects.get()
        resp = auth(self.directeur).post(f'{self.BASE}{element.pk}/restaurer/')
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.data['restaure'])
        self.lead.refresh_from_db()
        self.assertFalse(self.lead.is_archived)
        element.refresh_from_db()
        self.assertIsNotNone(element.restaure_le)
        # L'entrée quitte la corbeille active mais reste au journal.
        self.assertEqual(len(rows(auth(self.directeur).get(self.BASE))), 0)
        self.assertEqual(
            len(rows(auth(self.directeur).get(self.BASE, {'restaures': '1'}))), 1)

    def test_restaurer_refuse_a_un_commercial(self):
        archiver(self.lead)
        element = ElementSupprime.objects.get()
        resp = auth(self.commercial).post(f'{self.BASE}{element.pk}/restaurer/')
        self.assertEqual(resp.status_code, 403)

    def test_restaurer_refuse_une_seconde_fois(self):
        archiver(self.lead)
        element = ElementSupprime.objects.get()
        api = auth(self.directeur)
        self.assertEqual(api.post(f'{self.BASE}{element.pk}/restaurer/').status_code, 200)
        self.assertEqual(api.post(f'{self.BASE}{element.pk}/restaurer/').status_code, 400)

    def test_restaurer_refuse_une_entree_dune_autre_societe(self):
        archiver(self.lead)
        element = ElementSupprime.objects.get()
        resp = auth(self.directeur_b).post(f'{self.BASE}{element.pk}/restaurer/')
        self.assertEqual(resp.status_code, 404)

    def test_restaurer_creates_an_audit_entry(self):
        """NTUX38 — traçabilité audit (modèle `audit.AuditLog` existant)."""
        from apps.audit.models import AuditLog

        archiver(self.lead, type_libelle='Lead')
        element = ElementSupprime.objects.get()
        dernier = AuditLog.objects.order_by('-id').values_list(
            'id', flat=True).first() or 0
        resp = auth(self.directeur).post(f'{self.BASE}{element.pk}/restaurer/')
        self.assertEqual(resp.status_code, 200, resp.data)
        # Une restauration écrit DEUX traces distinctes et toutes deux
        # voulues : la trace CRUD automatique du Lead ré-enregistré
        # (``apps.audit.signals`` suit ``crm.Lead``) et, EN PLUS, la trace
        # NTUX38 de l'action corbeille. On verrouille donc celle qui appartient
        # à NTUX38 — jamais le total, qui interdirait toute autre trace légitime.
        nouvelles = list(AuditLog.objects.filter(id__gt=dernier))
        corbeille = [e for e in nouvelles
                     if 'corbeille' in (e.detail or '').lower()]
        self.assertEqual(len(corbeille), 1, [e.detail for e in nouvelles])
        entry = corbeille[0]
        self.assertEqual(entry.user, self.directeur)
        self.assertEqual(entry.company, self.co_a)


class NTUX24ExportXlsxTests(CorbeilleBase):
    """NTUX24 — export .xlsx du journal de corbeille (audit de rétention)."""

    def test_export_xlsx_forbidden_for_commercial(self):
        resp = auth(self.commercial).get(f'{self.BASE}export-xlsx/')
        self.assertEqual(resp.status_code, 403)

    def test_export_xlsx_reflects_the_same_filters_as_list(self):
        archiver(self.lead, type_libelle='Lead')
        autre = Lead.objects.create(company=self.co_a, nom='Bennani')
        archiver(autre, type_libelle='Devis')
        resp = auth(self.directeur).get(f'{self.BASE}export-xlsx/', {'type': 'Lead'})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        from openpyxl import load_workbook
        import io
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        self.assertEqual(ws.max_row, 2)  # en-tête + 1 seule ligne (filtre type=Lead)
        header = [c.value for c in ws[1]]
        self.assertEqual(
            header, ['Type', 'Libellé', 'Supprimé par', 'Supprimé le', 'Expire le', 'Statut'])
        self.assertEqual(ws.cell(row=2, column=1).value, 'Lead')
        self.assertEqual(ws.cell(row=2, column=6).value, 'Actif')

    def test_export_xlsx_marks_a_restored_row(self):
        archiver(self.lead, user=self.directeur, type_libelle='Lead')
        element = ElementSupprime.objects.get()
        auth(self.directeur).post(f'{self.BASE}{element.pk}/restaurer/')
        resp = auth(self.directeur).get(f'{self.BASE}export-xlsx/', {'restaures': '1'})
        from openpyxl import load_workbook
        import io
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        self.assertEqual(ws.cell(row=2, column=6).value, 'Restauré')

    def test_export_xlsx_excludes_restored_rows_by_default(self):
        archiver(self.lead, type_libelle='Lead')
        element = ElementSupprime.objects.get()
        auth(self.directeur).post(f'{self.BASE}{element.pk}/restaurer/')
        resp = auth(self.directeur).get(f'{self.BASE}export-xlsx/')
        from openpyxl import load_workbook
        import io
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        self.assertEqual(ws.max_row, 1)  # en-tête seul, la ligne restaurée est exclue


class NTUX26AvertissementRestaurationTests(CorbeilleBase):
    """NTUX26 — avertissement best-effort de l'assistant de restauration en
    masse (jamais bloquant, `None` sans signal exploitable)."""

    def test_no_signal_in_snapshot_means_no_warning(self):
        archiver(self.lead, donnees={'ville': 'Rabat'})
        resp = auth(self.directeur).get(self.BASE)
        self.assertIsNone(rows(resp)[0]['avertissement_restauration'])

    def test_responsable_disparu_declenche_un_avertissement(self):
        archiver(self.lead, donnees={'responsable_id': 999999})
        resp = auth(self.directeur).get(self.BASE)
        self.assertIn("n'existe plus", rows(resp)[0]['avertissement_restauration'])

    def test_responsable_desactive_declenche_un_avertissement(self):
        inactif = make_user(self.co_a, 'trash-inactif')
        inactif.is_active = False
        inactif.save(update_fields=['is_active'])
        archiver(self.lead, donnees={'responsable_id': inactif.id})
        resp = auth(self.directeur).get(self.BASE)
        self.assertIn('désactivé', rows(resp)[0]['avertissement_restauration'])

    def test_responsable_actif_ne_declenche_aucun_avertissement(self):
        archiver(self.lead, donnees={'responsable_id': self.directeur.id})
        resp = auth(self.directeur).get(self.BASE)
        self.assertIsNone(rows(resp)[0]['avertissement_restauration'])

    def test_element_deja_restaure_naffiche_jamais_davertissement(self):
        archiver(self.lead, donnees={'responsable_id': 999999})
        element = ElementSupprime.objects.get()
        auth(self.directeur).post(f'{self.BASE}{element.pk}/restaurer/')
        resp = auth(self.directeur).get(self.BASE, {'restaures': '1'})
        self.assertIsNone(rows(resp)[0]['avertissement_restauration'])


class NTUX31PermissionsFinesTests(CorbeilleBase):
    """NTUX31 — `ux.corbeille.consulter`/`ux.corbeille.restaurer` s'ajoutent EN
    PLUS du palier hérité `IsAdminOrResponsableTier`, sans jamais retirer
    l'accès d'un compte SANS rôle fin (repli légacy, voir `CorbeilleBase`)."""

    def setUp(self):
        super().setUp()
        # Rôle fin « responsable » (porte `users_voir` -> menu_tier=responsable,
        # donc passerait déjà IsAdminOrResponsableTier côté legacy) SANS les
        # deux codes NTUX31 : la garde fine doit désormais refuser.
        self.role_sans_code = Role.objects.create(
            company=self.co_a, nom='Responsable maison', est_systeme=False,
            permissions=['users_voir', 'crm_voir'],
        )
        self.role_avec_code = Role.objects.create(
            company=self.co_a, nom='Responsable outillé', est_systeme=False,
            permissions=[
                'users_voir', 'crm_voir',
                'ux_corbeille_consulter', 'ux_corbeille_restaurer',
            ],
        )
        self.user_sans_code = User.objects.create_user(
            username='trash31-sans', password='x', company=self.co_a,
            role=self.role_sans_code)
        self.user_avec_code = User.objects.create_user(
            username='trash31-avec', password='x', company=self.co_a,
            role=self.role_avec_code)

    def test_liste_refusee_sans_le_code_fin(self):
        archiver(self.lead)
        resp = auth(self.user_sans_code).get(self.BASE)
        self.assertEqual(resp.status_code, 403, resp.data)

    def test_liste_autorisee_avec_le_code_fin(self):
        archiver(self.lead)
        resp = auth(self.user_avec_code).get(self.BASE)
        self.assertEqual(resp.status_code, 200, resp.data)

    def test_restaurer_refuse_sans_le_code_fin(self):
        archiver(self.lead)
        element = ElementSupprime.objects.get()
        resp = auth(self.user_sans_code).post(f'{self.BASE}{element.pk}/restaurer/')
        self.assertEqual(resp.status_code, 403, resp.data)

    def test_restaurer_autorise_avec_le_code_fin(self):
        archiver(self.lead)
        element = ElementSupprime.objects.get()
        resp = auth(self.user_avec_code).post(f'{self.BASE}{element.pk}/restaurer/')
        self.assertEqual(resp.status_code, 200, resp.data)


class RestaurationTests(CorbeilleBase):
    def test_le_restaurateur_de_lapp_cible_a_la_priorite(self):
        appels = []

        def restaurateur_crm(element):
            appels.append(element.pk)
            lead = Lead.objects.get(pk=element.object_id)
            lead.is_archived = False
            lead.save(update_fields=['is_archived'])
            return lead

        registry.enregistrer_restaurateur('crm.lead', restaurateur_crm)
        archiver(self.lead)
        element = ElementSupprime.objects.get()
        obj = restaurer(element)
        self.assertEqual(appels, [element.pk])
        self.assertEqual(obj.pk, self.lead.pk)
        self.lead.refresh_from_db()
        self.assertFalse(self.lead.is_archived)

    def test_cible_disparue_ferme_lentree_sans_erreur(self):
        archiver(self.lead)
        element = ElementSupprime.objects.get()
        self.lead.delete()
        self.assertIsNone(restaurer(element))
        element.refresh_from_db()
        self.assertIsNotNone(element.restaure_le)

    def test_cible_sans_soft_delete_leve_une_erreur_explicite(self):
        element = ElementSupprime.objects.create(
            company=self.co_a,
            content_type=ContentType.objects.get_for_model(Company),
            object_id=self.co_a.pk, type_libelle='Société')
        with self.assertRaises(RestaurationImpossible):
            restaurer(element)


class PurgeTests(CorbeilleBase):
    def _perimer(self, element, jours=1):
        element.expire_le = timezone.now() - timedelta(days=jours)
        element.save(update_fields=['expire_le'])
        return element

    def test_purge_supprime_les_entrees_expirees_seulement(self):
        archiver(self.lead)
        perimee = self._perimer(ElementSupprime.objects.get())
        recente = Lead.objects.create(company=self.co_a, nom='Bennani')
        archiver(recente)
        purger_expires()
        restantes = list(ElementSupprime.objects.values_list('pk', flat=True))
        self.assertNotIn(perimee.pk, restantes)
        self.assertEqual(len(restantes), 1)

    def test_purge_ne_touche_jamais_la_cible(self):
        archiver(self.lead)
        self._perimer(ElementSupprime.objects.get())
        purger_expires()
        self.assertTrue(Lead.objects.filter(pk=self.lead.pk).exists())

    def test_commande_purger_corbeille(self):
        # `purger_corbeille_transverse` : `purger_corbeille` tout court est déjà
        # pris par GED25 et Django résoudrait le doublon vers `apps.ged`.
        archiver(self.lead)
        self._perimer(ElementSupprime.objects.get())
        call_command('purger_corbeille_transverse', '--dry-run')
        self.assertEqual(ElementSupprime.objects.count(), 1)
        call_command('purger_corbeille_transverse')
        self.assertEqual(ElementSupprime.objects.count(), 0)

    def test_tache_beat_purge_les_entrees_expirees(self):
        # NTUX29 — la tâche Celery planifiée (`trash.purger_corbeille_transverse`,
        # `erp_agentique/celery.py`) délègue au même service que la commande.
        from .tasks import purger_corbeille_transverse

        archiver(self.lead)
        self._perimer(ElementSupprime.objects.get())
        resultat = purger_corbeille_transverse()
        self.assertEqual(resultat, {self.co_a.id: 1})
        self.assertEqual(ElementSupprime.objects.count(), 0)

    def test_tache_beat_est_planifiee(self):
        # NTUX29 — garde-fou minimal : la tâche est bien enregistrée dans le
        # beat_schedule (le bug dominant du repo est une tâche bâtie/testée
        # mais jamais planifiée, cf. QX11).
        from erp_agentique.celery import app as celery_app

        taches = {v['task'] for v in celery_app.conf.beat_schedule.values()}
        self.assertIn('trash.purger_corbeille_transverse', taches)
