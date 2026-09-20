"""NTI18N46 — export XLSX du glossaire terminologique.

Le critère d'acceptation porte sur le FICHIER : il doit s'ouvrir sans corruption
et signaler visuellement les cases vides. Les tests relisent donc le classeur
produit avec ``openpyxl`` (ce que fait Excel/LibreOffice) au lieu de se fier au
code qui l'a écrit.

Run :
    python manage.py test apps.parametres.tests_nti18n46_glossaire_export -v 2
"""
import io

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import AccessToken

from apps.parametres.glossaire_export import (
    GLOSSAIRE_HEADERS,
    ROUGE_MANQUANT,
    construire_classeur,
    lignes_mentions,
    lignes_statuts,
    lignes_unites,
)
from apps.parametres.models_documents import DocumentTemplates
from apps.parametres.models_units import UniteMesure
from apps.roles.models import ALL_PERMISSIONS, DIRECTEUR_PERMISSIONS, Role
from authentication.models import Company
from core.i18n_content import set_translation

User = get_user_model()

EXPORT = '/api/django/parametres/traductions/glossaire-export/'


def _auth(user):
    api = APIClient()
    api.credentials(HTTP_AUTHORIZATION=f'Bearer {AccessToken.for_user(user)}')
    return api


def _relire(wb):
    """Sérialise puis relit le classeur — le chemin qu'Excel emprunte."""
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return load_workbook(buf)


class SourcesTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(
            nom='NTI18N46 Co', slug='nti18n46-co')

    def test_statuts_livrent_les_trois_langues(self):
        lignes = lignes_statuts(self.company)
        self.assertTrue(lignes)
        for domaine, cle, fr, en, ar in lignes:
            with self.subTest(domaine=domaine, cle=cle):
                self.assertTrue(fr)
                self.assertTrue(en)
                self.assertTrue(ar)

    def test_unite_sans_variante_laisse_en_et_ar_vides(self):
        UniteMesure.objects.create(
            company=self.company, code='m', libelle='Mètre', actif=True)
        (_dom, code, fr, en, ar), = lignes_unites(self.company)
        self.assertEqual(code, 'm')
        self.assertEqual(fr, 'Mètre')
        self.assertEqual(en, '')
        self.assertEqual(ar, '')

    def test_unite_avec_variante_remonte_la_traduction(self):
        unite = UniteMesure.objects.create(
            company=self.company, code='kg', libelle='Kilogramme', actif=True)
        set_translation(unite, 'libelle', 'en', 'Kilogram')
        (_dom, _code, _fr, en, ar), = lignes_unites(self.company)
        self.assertEqual(en, 'Kilogram')
        self.assertEqual(ar, '')

    def test_unite_inactive_exclue(self):
        UniteMesure.objects.create(
            company=self.company, code='lot', libelle='Lot', actif=False)
        self.assertEqual(lignes_unites(self.company), [])

    def test_mention_non_renseignee_nest_pas_listee(self):
        """« Rien à traduire » n'est pas « traduction manquante »."""
        DocumentTemplates.objects.create(company=self.company)
        self.assertEqual(lignes_mentions(self.company), [])

    def test_mention_renseignee_listee_avec_ses_variantes(self):
        gabarit = DocumentTemplates.objects.create(
            company=self.company, cgv_titre='Conditions générales',
            cgv_bullets=['Acompte de 40 %', ''])
        set_translation(gabarit, 'cgv_titre', 'ar', 'الشروط العامة')

        lignes = lignes_mentions(self.company)

        par_champ = {ligne[1]: ligne for ligne in lignes}
        self.assertIn('cgv_titre', par_champ)
        self.assertEqual(par_champ['cgv_titre'][3], '')
        self.assertEqual(par_champ['cgv_titre'][4], 'الشروط العامة')
        # La puce vide est ignorée, la première est listée.
        self.assertIn('cgv_bullets.1', par_champ)
        self.assertNotIn('cgv_bullets.2', par_champ)

    def test_societe_absente_ne_leve_pas(self):
        self.assertEqual(lignes_unites(None), [])
        self.assertEqual(lignes_mentions(None), [])


class ClasseurTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(
            nom='NTI18N46 Classeur', slug='nti18n46-classeur')
        UniteMesure.objects.create(
            company=self.company, code='m', libelle='Mètre', actif=True)

    def test_trois_onglets_nommes_par_domaine(self):
        wb = _relire(construire_classeur(self.company))
        self.assertEqual(wb.sheetnames,
                         ['Statuts', 'Unités', 'Mentions légales'])

    def test_chaque_onglet_porte_la_ligne_den_tetes(self):
        wb = _relire(construire_classeur(self.company))
        for nom in wb.sheetnames:
            with self.subTest(onglet=nom):
                entetes = [c.value for c in wb[nom][1]]
                self.assertEqual(entetes, GLOSSAIRE_HEADERS)

    def test_case_manquante_peinte_en_rouge(self):
        wb = _relire(construire_classeur(self.company))
        ws = wb['Unités']
        for colonne in ('D', 'E'):
            with self.subTest(colonne=colonne):
                cellule = ws[f'{colonne}2']
                # openpyxl relit une chaîne vide en ``None`` : on teste
                # « pas de contenu », pas une représentation précise.
                self.assertFalse(str(cellule.value or '').strip())
                self.assertIn(ROUGE_MANQUANT,
                              str(cellule.fill.start_color.rgb or ''))

    def test_case_remplie_nest_pas_peinte(self):
        wb = _relire(construire_classeur(self.company))
        ws = wb['Statuts']
        cellule = ws['D2']
        self.assertTrue(cellule.value)
        self.assertNotIn(ROUGE_MANQUANT,
                         str(cellule.fill.start_color.rgb or ''))

    def test_texte_a_risque_neutralise(self):
        """Un texte commençant par « = » ne doit jamais devenir une formule."""
        DocumentTemplates.objects.create(
            company=self.company, bpa_titre='=HYPERLINK("http://x","Voir")')
        wb = _relire(construire_classeur(self.company))
        valeurs = [
            ws_row[2].value
            for ws_row in wb['Mentions légales'].iter_rows(min_row=2)
        ]
        self.assertTrue(any(str(v).startswith("'=") for v in valeurs), valeurs)


class EndpointTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(
            nom='NTI18N46 Endpoint', slug='nti18n46-endpoint')
        UniteMesure.objects.create(
            company=self.company, code='m', libelle='Mètre', actif=True)
        self.admin = User.objects.create_user(
            username='nti18n46-admin', password='x', role_legacy='admin',
            company=self.company)

    def test_export_telechargeable_et_relisible(self):
        resp = _auth(self.admin).get(EXPORT)
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertIn('spreadsheetml', resp['Content-Type'])
        self.assertIn('glossaire-traductions.xlsx',
                      resp['Content-Disposition'])
        wb = load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.sheetnames,
                         ['Statuts', 'Unités', 'Mentions légales'])

    def test_export_refuse_sans_localisation_gerer(self):
        role = Role.objects.create(
            company=self.company, nom='Sans localisation',
            permissions=[c for c in DIRECTEUR_PERMISSIONS
                         if c != 'localisation_gerer'])
        user = User.objects.create_user(
            username='nti18n46-sans', password='x', role_legacy='admin',
            company=self.company, role=role)
        resp = _auth(user).get(EXPORT)
        self.assertEqual(resp.status_code, 403, resp.content)
        self.assertIn('localisation_gerer', ALL_PERMISSIONS)
