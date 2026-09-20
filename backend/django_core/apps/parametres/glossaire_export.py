"""NTI18N46 — export XLSX du glossaire terminologique pour relecture métier.

CE QUE ÇA RÉSOUT : le glossaire (NTI18N25) vit dans du code et dans des tables ;
un relecteur métier — le commercial arabophone, le comptable sénégalais — ne peut
pas l'auditer depuis l'ERP, et personne ne voit d'un coup d'œil CE QUI MANQUE.
Ce module rend un classeur à onglets, un par domaine, où chaque case de
traduction absente est PEINTE EN ROUGE : le relecteur ouvre, trie, annote hors
ligne.

TROIS ONGLETS, TROIS SOURCES RÉELLES (aucune n'est inventée ici) :
  * « Statuts » — le catalogue statique ``i18n_labels.STATUT_LABELS`` (fr/en/ar
    livrés), avec la colonne FR passée par ``selectors.statut_libelle`` pour que
    le relecteur lise la formulation de SA société (surcharge ``StatutConfig``,
    qui reste FR par construction) ;
  * « Unités » — le référentiel société ``UniteMesure`` (libellé FR saisi), ses
    variantes EN/AR lues dans ``core.ContentTranslation`` (YHARD4, le mécanisme
    de traduction des données SAISIES) ;
  * « Mentions légales » — les textes éditables ``DocumentTemplates`` que la
    société a RÉELLEMENT renseignés, mêmes variantes EN/AR.

Seules les lignes dont le FR existe sont listées pour les deux derniers onglets :
« rien à traduire » et « traduction manquante » sont deux états distincts, et
peindre en rouge un texte que la société n'a jamais écrit serait un faux signal.

ROUGE PEINT, PAS UNE RÈGLE CONDITIONNELLE : la condition (« cette variante est
absente ») est connue à la génération, donc la couleur est posée cellule par
cellule. Une règle de mise en forme conditionnelle openpyxl exige de surcroît
``bgColor`` en plus de ``fgColor`` pour être visible — un remplissage invisible
est exactement le genre de signal muet que cette tâche existe pour supprimer.
"""
from __future__ import annotations

from apps.records.xlsx import XLSX_CONTENT_TYPE, coerce_cell, neutralize_cell

#: Langues de la relecture : le FR est la source, EN et AR sont les cibles
#: contrôlées (mêmes trois langues que le reste du cadre i18n léger).
LANGUES_CIBLES = ('en', 'ar')

#: En-têtes communs aux trois onglets.
GLOSSAIRE_HEADERS = ['Domaine', 'Clé', 'FR', 'EN', 'AR']

#: Rouge pâle lisible en noir, posé sur une case de traduction absente.
ROUGE_MANQUANT = 'FFC7CE'

#: Champs texte de ``DocumentTemplates`` qui portent une mention légale ou
#: contractuelle destinée au client — donc relus/traduits. ``cgv_bullets``
#: (liste JSON) est aplati en une ligne par puce.
CHAMPS_MENTIONS = (
    ('validite_badge_p1', "Validité (badge page 1)"),
    ('validite_onepage', "Validité (une page)"),
    ('cgv_titre', "Conditions générales — titre"),
    ('garantie_titre', "Garantie — titre"),
    ('garantie_detail', "Garantie — détail"),
    ('garantie_perf_label', "Garantie — performance"),
    ('bpa_titre', "Bon pour accord — titre"),
    ('bpa_mention', "Bon pour accord — mention"),
    ('acceptance_stamp', "Tampon d'acceptation"),
)


def _variantes(instance):
    """``{locale: {champ: valeur}}`` des traductions de contenu d'un objet.

    Passe par ``core.i18n_content.translations_for`` (jamais une requête
    ``ContentTranslation`` réécrite ici). Dict vide — jamais d'exception — pour
    un objet sans société ou non enregistré.
    """
    from core.i18n_content import translations_for

    try:
        return translations_for(instance)
    except Exception:  # noqa: BLE001 — un export ne casse jamais sur ce point
        return {}


def lignes_statuts(company):
    """Une ligne par (domaine métier, clé de statut) du catalogue NTI18N25."""
    from .i18n_labels import STATUT_LABELS, statut_label
    from .selectors import statut_libelle

    lignes = []
    for domaine, par_cle in STATUT_LABELS.items():
        for cle in par_cle:
            lignes.append([
                domaine,
                cle,
                statut_libelle(company, domaine, cle, 'fr'),
                statut_label(domaine, cle, 'en'),
                statut_label(domaine, cle, 'ar'),
            ])
    return lignes


def lignes_unites(company):
    """Une ligne par unité de mesure ACTIVE du référentiel de la société."""
    from .models_units import UniteMesure

    if company is None:
        return []
    lignes = []
    for unite in UniteMesure.objects.filter(company=company, actif=True):
        variantes = _variantes(unite)
        lignes.append([
            'unite',
            unite.code,
            unite.libelle,
            (variantes.get('en') or {}).get('libelle', ''),
            (variantes.get('ar') or {}).get('libelle', ''),
        ])
    return lignes


def lignes_mentions(company):
    """Une ligne par mention légale que la société a RÉELLEMENT renseignée."""
    from .models_documents import DocumentTemplates

    if company is None:
        return []
    gabarit = DocumentTemplates.objects.filter(company=company).first()
    if gabarit is None:
        return []
    variantes = _variantes(gabarit)
    lignes = []
    for champ, libelle in CHAMPS_MENTIONS:
        valeur = (getattr(gabarit, champ, '') or '').strip()
        if not valeur:
            continue
        lignes.append([
            libelle,
            champ,
            valeur,
            (variantes.get('en') or {}).get(champ, ''),
            (variantes.get('ar') or {}).get(champ, ''),
        ])
    puces = gabarit.cgv_bullets or []
    for index, puce in enumerate(puces, start=1):
        texte = (str(puce) or '').strip()
        if not texte:
            continue
        champ = f'cgv_bullets.{index}'
        lignes.append([
            "Conditions générales — puce",
            champ,
            texte,
            (variantes.get('en') or {}).get(champ, ''),
            (variantes.get('ar') or {}).get(champ, ''),
        ])
    return lignes


#: Onglets du classeur, dans l'ordre : nom d'onglet -> fonction de lignes.
ONGLETS = (
    ('Statuts', lignes_statuts),
    ('Unités', lignes_unites),
    ('Mentions légales', lignes_mentions),
)


def _ecrire_onglet(ws, titre, lignes, rouge, gras):
    ws.title = titre
    ws.append(list(GLOSSAIRE_HEADERS))
    for cellule in ws[1]:
        cellule.font = gras
    largeurs = [len(entete) for entete in GLOSSAIRE_HEADERS]
    for ligne in lignes:
        cellules = [coerce_cell(neutralize_cell(v)) for v in ligne]
        ws.append(cellules)
        for index, valeur in enumerate(cellules):
            if index < len(largeurs):
                largeurs[index] = max(largeurs[index], len(str(valeur)))
        # Colonnes EN (4) et AR (5) : rouge quand la variante manque.
        for colonne in (4, 5):
            if not str(cellules[colonne - 1] or '').strip():
                ws.cell(row=ws.max_row, column=colonne).fill = rouge
    for index in range(len(GLOSSAIRE_HEADERS)):
        lettre = ws.cell(row=1, column=index + 1).column_letter
        ws.column_dimensions[lettre].width = min(
            max(largeurs[index] + 2, 12), 60)


def construire_classeur(company):
    """Classeur ``openpyxl`` du glossaire d'une société (un onglet par domaine).

    ``openpyxl`` est une dépendance pré-approuvée ; import local pour ne la
    charger qu'au moment d'un export.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    rouge = PatternFill(
        fill_type='solid', start_color=ROUGE_MANQUANT,
        end_color=ROUGE_MANQUANT)
    gras = Font(bold=True)

    wb = Workbook()
    for position, (titre, source) in enumerate(ONGLETS):
        ws = wb.active if position == 0 else wb.create_sheet()
        _ecrire_onglet(ws, titre, source(company), rouge, gras)
    return wb


def reponse_export(company, filename='glossaire-traductions.xlsx'):
    """Réponse HTTP téléchargeable du classeur de ``company``."""
    from django.http import HttpResponse

    response = HttpResponse(content_type=XLSX_CONTENT_TYPE)
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    construire_classeur(company).save(response)
    return response
