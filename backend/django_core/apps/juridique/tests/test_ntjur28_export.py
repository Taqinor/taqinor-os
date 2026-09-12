"""NTJUR28 — export .xlsx du registre des dossiers (assureur / due diligence).

Critère d'acceptation : « l'export xlsx s'ouvre sans erreur et exclut les
dossiers confidentiels pour un exportateur non autorisé ».
"""
import io
from datetime import date
from decimal import Decimal

from django.test import TestCase

from apps.juridique.models import DossierJuridique

from ._base import auth, make_admin, make_company, make_responsable

URL = '/api/django/juridique/dossiers/export/'
XLSX = ('application/vnd.openxmlformats-officedocument.'
        'spreadsheetml.sheet')


def _lignes(reponse):
    """Ouvre réellement le classeur et rend ses lignes (en-tête inclus)."""
    from openpyxl import load_workbook

    classeur = load_workbook(io.BytesIO(reponse.content))
    return [list(r) for r in classeur.active.iter_rows(values_only=True)]


class ExportDossiersTests(TestCase):
    def setUp(self):
        self.company = make_company('jur-x28-co', 'Juridique X28')
        self.autre = make_company('jur-x28-autre', 'Juridique X28 Autre')
        self.admin = make_admin(self.company, 'jur-x28-admin')
        self.api = auth(self.admin)
        self.public = DossierJuridique.objects.create(
            company=self.company, reference='JUR-2026-0001',
            titre='Litige fournisseur', date_ouverture=date(2026, 2, 1),
            montant_en_jeu=Decimal('120000'))
        self.secret = DossierJuridique.objects.create(
            company=self.company, reference='JUR-2026-0002',
            titre='Affaire direction', date_ouverture=date(2026, 2, 2),
            confidentialite=(
                DossierJuridique.NiveauConfidentialite.CONFIDENTIEL))
        DossierJuridique.objects.create(
            company=self.autre, reference='JUR-2026-9001',
            titre='Hors société', date_ouverture=date(2026, 2, 3))

    def test_le_classeur_s_ouvre_et_porte_les_bonnes_colonnes(self):
        resp = self.api.get(URL)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], XLSX)
        lignes = _lignes(resp)
        self.assertEqual(lignes[0][0], 'Référence')
        self.assertIn('Provision comptabilisée', lignes[0])
        # Jamais de donnée hors-périmètre dans l'en-tête.
        entete = ' '.join(str(c) for c in lignes[0]).lower()
        self.assertNotIn('achat', entete)
        self.assertNotIn('marge', entete)

    def test_l_export_est_scope_societe(self):
        references = {r[0] for r in _lignes(self.api.get(URL))[1:]}
        self.assertNotIn('JUR-2026-9001', references)

    def test_un_exportateur_non_autorise_n_a_pas_les_confidentiels(self):
        responsable = make_responsable(self.company, 'jur-x28-resp')
        resp = auth(responsable).get(URL)
        self.assertEqual(resp.status_code, 200)
        references = {r[0] for r in _lignes(resp)[1:]}
        self.assertIn('JUR-2026-0001', references)
        self.assertNotIn('JUR-2026-0002', references)
        # L'administrateur, lui, les a.
        refs_admin = {r[0] for r in _lignes(self.api.get(URL))[1:]}
        self.assertIn('JUR-2026-0002', refs_admin)

    def test_filtres_statut_et_periode(self):
        par_statut = self.api.get(f'{URL}?statut=ouvert')
        self.assertEqual(par_statut.status_code, 200)
        self.assertEqual(len(_lignes(par_statut)) - 1, 2)

        par_annee = self.api.get(f'{URL}?periode=2025')
        self.assertEqual(len(_lignes(par_annee)) - 1, 0)

        invalide = self.api.get(f'{URL}?periode=abcd')
        self.assertEqual(invalide.status_code, 400)
        self.assertIn('periode', invalide.data)
