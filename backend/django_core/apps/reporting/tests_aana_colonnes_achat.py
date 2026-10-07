"""AANA25 (C-AANA-016) — aucune colonne à prix d'achat ne sort d'un rapport :
rapport planifié (requête et tableau de bord), lien public partagé et export
du générateur ; l'export s'exécute au nom du LECTEUR.

Scénarios R3/R4 : une requête enregistrée ``stock_produits`` [nom,
prix_achat] d'un admin, et une définition partagée [nom, valeur_achat].
Avant le correctif : en-tête ['nom', 'prix_achat'] dans le rendu planifié ;
export CSV par un lecteur sans ``prix_achat_voir`` = 200 avec
``nom,valeur_achat`` = 300.00.

La liste des colonnes interdites est DÉRIVÉE des ``gated_fields`` des
datasets (``core.data_explorer``) — aucun mock : vrais datasets enregistrés
par ``apps.stock``, vraie base, vrai rendu xlsx relu par openpyxl.
"""
import csv
import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import AccessToken

from apps.reporting import scheduled_reports
from apps.reporting.diffusion_views import lien_rapport
from apps.reporting.models import RapportDefinition, SavedReport
from apps.reporting.rapport_builder import champs_gated
from apps.roles.models import Role
from apps.stock.models import Produit
from authentication.models import Company
from core.models import Dashboard, SavedQuery

User = get_user_model()


def _entetes_xlsx(contenu):
    from openpyxl import load_workbook
    feuille = load_workbook(io.BytesIO(contenu)).active
    return [c.value for c in next(feuille.iter_rows(max_row=1))]


class TestColonnesAchat(TestCase):
    def setUp(self):
        self.company = Company.objects.get_or_create(
            slug='aana25-co', defaults={'nom': 'AANA25 Co'})[0]
        self.admin = User.objects.create_user(
            username='aana25_admin', password='x', role_legacy='admin',
            company=self.company)
        self.assertTrue(self.admin.can_view_buy_prices)
        Produit.objects.create(
            company=self.company, nom='Panneau', sku='AANA25-P',
            prix_vente=Decimal('150'), prix_achat=Decimal('100'),
            quantite_stock=3)

    def test_liste_derivee_des_gated_fields(self):
        self.assertTrue({'prix_achat', 'valeur_achat'}
                        <= champs_gated('stock_produits'))
        self.assertIn('montant', champs_gated('stock_bcf'))

    def test_rapport_planifie_sans_prix_achat(self):
        requete = SavedQuery.objects.create(
            company=self.company, owner=self.admin, titre='Catalogue achat',
            dataset='stock_produits', spec={'select': ['nom', 'prix_achat']})
        rapport = SavedReport.objects.create(
            company=self.company, owner=self.admin, name='Catalogue',
            target_kind=SavedReport.TargetKind.QUERY, cible_id=requete.pk,
            recipients='externe@example.com')

        contenu, _titre, nom, _mime = scheduled_reports.rendre_rapport(rapport)
        self.assertTrue(nom.endswith('.xlsx'))
        self.assertEqual(_entetes_xlsx(contenu), ['nom'])

        # Lien PUBLIC partagé : même rendu, même filtre.
        reponse = APIClient().get(lien_rapport(rapport))
        self.assertEqual(reponse.status_code, 200)
        self.assertEqual(_entetes_xlsx(reponse.content), ['nom'])

    def test_rapport_planifie_projection_par_defaut(self):
        """Sans ``select``, la projection par défaut d'un admin ne fait pas
        sortir les champs sous permission."""
        requete = SavedQuery.objects.create(
            company=self.company, owner=self.admin, titre='Tout',
            dataset='stock_produits', spec={})
        rapport = SavedReport.objects.create(
            company=self.company, owner=self.admin, name='Tout',
            target_kind=SavedReport.TargetKind.QUERY, cible_id=requete.pk)
        contenu, _t, _n, _m = scheduled_reports.rendre_rapport(rapport)
        entetes = _entetes_xlsx(contenu)
        self.assertIn('nom', entetes)
        self.assertNotIn('prix_achat', entetes)
        self.assertNotIn('valeur_achat', entetes)

    def test_dashboard_planifie_sans_valeur_achat(self):
        dashboard = Dashboard.objects.create(
            company=self.company, owner=self.admin, titre='Stock',
            layout={'widgets': [{
                'id': 'stock', 'titre': 'Stock', 'dataset': 'stock_produits',
                'spec': {'select': ['nom', 'valeur_achat']},
            }, {
                'id': 'somme', 'titre': 'Somme achat',
                'dataset': 'stock_produits',
                'spec': {'aggregates': [{'alias': 'total', 'fn': 'sum',
                                         'field': 'valeur_achat'}]},
            }]})
        rapport = SavedReport.objects.create(
            company=self.company, owner=self.admin, name='Stock',
            target_kind=SavedReport.TargetKind.DASHBOARD,
            cible_id=dashboard.pk)
        html = scheduled_reports.rendre_dashboard_html(rapport, dashboard)
        self.assertIn('Panneau', html)
        self.assertNotIn('valeur_achat', html)
        self.assertNotIn('prix_achat', html)

    def test_export_definition_lecteur(self):
        definition = RapportDefinition.objects.create(
            company=self.company, owner=self.admin, titre='Valeur stock',
            dataset='stock_produits', spec={'select': ['nom', 'valeur_achat']},
            partage=RapportDefinition.Partage.SOCIETE)
        role = Role.objects.create(
            company=self.company, nom='Lecteur stock',
            permissions=['stock_voir'])
        lecteur = User.objects.create_user(
            username='aana25_lecteur', password='x', company=self.company,
            role=role)
        self.assertFalse(lecteur.can_view_buy_prices)
        api = APIClient()
        api.credentials(
            HTTP_AUTHORIZATION=f'Bearer {AccessToken.for_user(lecteur)}')

        reponse = api.get(
            f'/api/django/reporting/rapport-definitions/{definition.pk}'
            '/export/?format=csv')
        self.assertEqual(reponse.status_code, 200)
        lignes = list(csv.reader(io.StringIO(
            reponse.content.decode('utf-8-sig'))))
        self.assertEqual(lignes[0], ['nom'])
        self.assertNotIn('300.00', {c for ligne in lignes for c in ligne})
