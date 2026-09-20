"""CAL179 — modules, chaînes et nomenclature : les quantités, et AUCUN prix.

Essais PURS (ni base, ni WeasyPrint) : ``openpyxl`` est déjà une dépendance du
dépôt, et le classeur produit est RELU cellule par cellule — un fichier « écrit
sans erreur » peut être vide.

Ce qui est prouvé :

* les quantités exportées sont celles de la géométrie stockée (pas une
  approximation, pas un compte recalculé) ;
* la rangée est un GROUPEMENT des ordonnées relevées : deux modules de même
  ordonnée sont sur la même rangée, deux modules d'ordonnées différentes non ;
* AUCUNE colonne, AUCUNE cellule ne porte de prix — et la garde refuse
  l'export si l'on en glisse un, plutôt que de le filtrer en silence ;
* l'injection de formules est neutralisée (la règle de ``apps/records/xlsx.py``
  n'est pas recodée, elle est appelée).

Run :
    python manage.py test apps.calepinage.tests.test_cal179_export_tableur -v2
"""
import copy
import io
import json
import pathlib

from django.test import SimpleTestCase

from apps.calepinage.services.export_tableur import (
    FEUILLES, ExportRefuse, classeur_octets, csv_octets, table_chaines,
    table_modules, table_nomenclature, tables_du_resultat,
    verifier_absence_de_prix,
)
from apps.calepinage.services.planche import geometrie_de_planche

from .test_cal171_planche import LAYOUT

RACINE_APP = pathlib.Path(__file__).resolve().parents[1]
RESULTAT = json.loads(
    (RACINE_APP / 'contract_samples' / 'calepinage_resultat.json')
    .read_text(encoding='utf-8'))['exemple']


def geometrie():
    return geometrie_de_planche(LAYOUT)


def relire(octets):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(octets))


class TableModulesTest(SimpleTestCase):
    def test_une_ligne_par_module_pose(self):
        _entetes, lignes = table_modules(geometrie(), RESULTAT)
        poses = sum(len(pan['modules']) for pan in geometrie()['pans'])
        self.assertEqual(len(lignes), poses)
        self.assertEqual(poses, 2)

    def test_la_position_relevee_est_exportee_en_metres(self):
        entetes, lignes = table_modules(geometrie(), RESULTAT)
        est = entetes.index('Est (m)')
        self.assertAlmostEqual(lignes[1][est] - lignes[0][est], 2.5, places=2)

    def test_deux_modules_de_meme_ordonnee_partagent_leur_rangee(self):
        entetes, lignes = table_modules(geometrie(), RESULTAT)
        rangee = entetes.index('Rangée')
        self.assertEqual(lignes[0][rangee], lignes[1][rangee])

    def test_un_module_d_une_autre_ordonnee_change_de_rangee(self):
        layout = copy.deepcopy(LAYOUT)
        layout['zones'][0]['geometry']['panels'][1]['cy'] = 4.0
        entetes, lignes = table_modules(geometrie_de_planche(layout),
                                        RESULTAT)
        rangee = entetes.index('Rangée')
        self.assertNotEqual(lignes[0][rangee], lignes[1][rangee])

    def test_une_orientation_absente_reste_vide_jamais_zero(self):
        layout = copy.deepcopy(LAYOUT)
        layout['zones'][0]['geometry'].pop('azimuthDeg')
        layout['zones'][0]['geometry'].pop('tiltDeg')
        entetes, lignes = table_modules(geometrie_de_planche(layout), None)
        self.assertIsNone(lignes[0][entetes.index('Azimut (°)')])
        self.assertIsNone(lignes[0][entetes.index('Inclinaison (°)')])


class TablesChainesEtNomenclatureTest(SimpleTestCase):
    def test_le_chainage_est_celui_publie_par_le_moteur(self):
        _entetes, lignes = table_chaines(RESULTAT)
        valeurs = dict((ligne[0], ligne[1]) for ligne in lignes)
        chainage = RESULTAT['electrique']['chainage']
        self.assertEqual(valeurs['Nombre de chaînes'], chainage['chaines'])
        self.assertEqual(valeurs['Modules par chaîne'],
                         chainage['modules_par_chaine'])

    def test_la_nomenclature_porte_designation_et_quantite(self):
        entetes, lignes = table_nomenclature(RESULTAT)
        self.assertEqual(entetes, ['Désignation', 'Quantité', 'Unité'])
        self.assertEqual(lignes[0][1], RESULTAT['pose']['total_modules'])
        self.assertIn('720', lignes[0][0])

    def test_sans_resultat_les_tables_restent_lisibles(self):
        for entetes, lignes in (table_chaines(None), table_nomenclature(None)):
            self.assertTrue(entetes)
            self.assertIsInstance(lignes, list)


class AucunPrixTest(SimpleTestCase):
    def test_aucune_colonne_de_prix_dans_les_trois_feuilles(self):
        for _titre, entetes, lignes in tables_du_resultat(geometrie(),
                                                          RESULTAT):
            for entete in entetes:
                self.assertNotIn('prix', str(entete).lower())
                self.assertNotIn('achat', str(entete).lower())
            for ligne in lignes:
                for cellule in ligne:
                    self.assertNotIn('prix_achat', str(cellule).lower())

    def test_la_garde_refuse_un_prix_glisse_dans_un_en_tete(self):
        with self.assertRaises(ExportRefuse) as capture:
            verifier_absence_de_prix(['Désignation', "Prix d'achat"], [])
        self.assertIn('prix', str(capture.exception).lower())

    def test_la_garde_refuse_un_prix_glisse_dans_une_cellule(self):
        with self.assertRaises(ExportRefuse):
            verifier_absence_de_prix(['Désignation'],
                                     [['Module 720 Wc — 1 200 MAD']])


class ClasseurTest(SimpleTestCase):
    def setUp(self):
        self.classeur = relire(
            classeur_octets(tables_du_resultat(geometrie(), RESULTAT)))

    def test_les_trois_feuilles_sont_la_et_nommees(self):
        self.assertEqual(self.classeur.sheetnames, list(FEUILLES))

    def test_la_feuille_modules_porte_ses_lignes(self):
        feuille = self.classeur[FEUILLES[0]]
        self.assertEqual(feuille.max_row, 3)  # en-tête + 2 modules
        self.assertEqual(feuille.cell(row=1, column=1).value, 'Pan')

    def test_l_injection_de_formule_est_neutralisee(self):
        tables = [('Modules', ['Pan'], [['=1+1']])]
        feuille = relire(classeur_octets(tables))['Modules']
        self.assertEqual(feuille.cell(row=2, column=1).value, "'=1+1")


class CsvTest(SimpleTestCase):
    def test_le_csv_reprend_la_feuille_demandee(self):
        octets = csv_octets(tables_du_resultat(geometrie(), RESULTAT),
                            feuille='Nomenclature')
        texte = octets.decode('utf-8-sig')
        self.assertTrue(texte.startswith('Désignation;Quantité;Unité'))
        self.assertIn('720 Wc;12;u', texte)

    def test_le_csv_s_ouvre_sans_ecran_d_import(self):
        # BOM utf-8 + point-virgule : ce qu'Excel fr-MA ouvre directement.
        octets = csv_octets(tables_du_resultat(geometrie(), RESULTAT))
        self.assertTrue(octets.startswith(b'\xef\xbb\xbf'))
        self.assertIn(b';', octets)

    def test_une_feuille_inconnue_est_refusee_en_la_nommant(self):
        with self.assertRaises(ExportRefuse) as capture:
            csv_octets(tables_du_resultat(geometrie(), RESULTAT),
                       feuille='Tarifs')
        self.assertIn('Tarifs', str(capture.exception))
