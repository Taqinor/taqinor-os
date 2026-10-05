"""ACAL295 — la vue restreinte au responsable est UN prédicat d'accès.

Constat C-ACAL-014 (audit 2026-10-04) : la restriction ne filtrait que la
LISTE et le détail ; comparer-projets, comparatif.xlsx, archiver,
restaurer-corbeille, depuis-lead et les portes modèle lisaient
``company + pk`` sans elle — un concepteur restreint comparait, archivait
ou ouvrait le calepinage d'un autre.

Tenu ici avec le réglage ÉCRIT par ``enregistrer_parametres`` (jamais un
mock), deux comptes réels de la même société : U (gérer, sans approuver) et
le Directeur (approbateur, responsable de K).

Run :
    python manage.py test apps.calepinage.tests.test_acal_vue_restreinte_routes -v2
"""
from __future__ import annotations

import io

from django.contrib.auth import get_user_model

from apps.calepinage.models import Calepinage
from apps.calepinage.permissions import CAL_GERER, CAL_VOIR
from apps.calepinage.services.archivage import archiver, est_archive
from apps.calepinage.services.comparaison_projets import MOTIF_INTROUVABLE
from apps.calepinage.services.creation import MESSAGE_HORS_VUE
from apps.calepinage.services.parametres import enregistrer_parametres
from apps.roles.models import Role
from apps.trash.selectors import ids_dans_corbeille
from apps.ventes.models import Devis

from .test_api_liste import URL, BaseApiCalepinage, url_detail

User = get_user_model()


class VueRestreinteRoutesTest(BaseApiCalepinage):
    def setUp(self):
        super().setUp()
        role = Role.objects.create(company=self.company, nom='Concepteur295',
                                   permissions=[CAL_VOIR, CAL_GERER])
        self.u = User.objects.create_user(
            username='acal295_u', password='x', company=self.company,
            role=role)
        self.api_u = self._client(self.u)
        self.k = Calepinage.objects.create(
            company=self.company, lead_id=self.lead.pk, titre='Toit de K',
            responsable=self.user, cree_par=self.user)
        self.m = Calepinage.objects.create(
            company=self.company, lead_id=self.lead_2.pk, titre='Toit de M',
            cree_par=self.u)
        enregistrer_parametres(self.company, {
            'presets': {'vue_restreinte_au_responsable': True}})

    def test_comparer_projets_met_un_calepinage_hors_vue_en_refus(self):
        reponse = self.api_u.post(f'{URL}comparer-projets/',
                                  {'ids': [self.k.pk, self.m.pk, 999999]},
                                  format='json')
        self.assertEqual(reponse.status_code, 200, reponse.data)
        self.assertEqual([ligne['id'] for ligne in reponse.data['lignes']],
                         [self.m.pk])
        refus = {r['id']: r['motif'] for r in reponse.data['refus']}
        self.assertEqual(refus, {self.k.pk: MOTIF_INTROUVABLE,
                                 999999: MOTIF_INTROUVABLE})

    def test_comparatif_xlsx_omet_hors_vue(self):
        from openpyxl import load_workbook

        reponse = self.api_u.get(
            f'{url_detail(self.m.pk)}comparatif.xlsx/', {'ids': self.k.pk})
        self.assertEqual(reponse.status_code, 200)
        contenu = b''.join(reponse.streaming_content) if getattr(
            reponse, 'streaming', False) else reponse.content
        classeur = load_workbook(io.BytesIO(contenu))
        textes = {str(cellule.value) for feuille in classeur.worksheets
                  for ligne in feuille.iter_rows() for cellule in ligne
                  if cellule.value is not None}
        self.assertTrue(any('Toit de M' in t for t in textes))
        self.assertFalse(any('Toit de K' in t for t in textes))

    def test_archiver_et_restaurer_hors_vue_404(self):
        avant = set(ids_dans_corbeille('calepinage.calepinage'))
        reponse = self.api_u.post(f'{url_detail(self.k.pk)}archiver/')
        self.assertEqual(reponse.status_code, 404, reponse.data)
        self.assertEqual(reponse.data, {'detail': 'Calepinage introuvable.'})
        self.assertEqual(set(ids_dans_corbeille('calepinage.calepinage')),
                         avant)
        self.assertEqual(self.api.get(url_detail(self.k.pk)).status_code, 200)

        archiver(self.k, user=self.user)
        reponse = self.api_u.post(
            f'{url_detail(self.k.pk)}restaurer-corbeille/')
        self.assertEqual(reponse.status_code, 404, reponse.data)
        self.assertEqual(reponse.data, {'detail': 'Calepinage introuvable.'})
        self.assertTrue(est_archive(self.k))

    def test_depuis_lead_hors_vue_409_sans_reference(self):
        avant = Calepinage.objects.filter(lead_id=self.lead.pk).count()
        reponse = self.api_u.post(f'{URL}depuis-lead/',
                                  {'lead': self.lead.pk}, format='json')
        self.assertEqual(reponse.status_code, 409, reponse.data)
        self.assertEqual(reponse.data, {'detail': MESSAGE_HORS_VUE})
        self.assertNotIn(str(self.k.pk), str(reponse.data))
        self.assertNotIn('Toit de K', str(reponse.data))
        self.assertEqual(
            Calepinage.objects.filter(lead_id=self.lead.pk).count(), avant)

    def test_depuis_modele_non_modele_404_et_devis_hors_vue_409(self):
        for url, corps in ((f'{URL}depuis-modele/',
                            {'modele_id': self.k.pk,
                             'lead_id': self.lead_2.pk}),
                           (f'{URL}depuis-modele/',
                            {'modele_id': self.m.pk,
                             'lead_id': self.lead_2.pk}),
                           (f'{URL}creer-depuis-modele/',
                            {'modele': self.k.pk, 'lead': self.lead_2.pk})):
            with self.subTest(url=url, corps=corps):
                reponse = self.api_u.post(url, corps, format='json')
                self.assertEqual(reponse.status_code, 404, reponse.data)
                self.assertIn('Modèle introuvable.', str(reponse.data))

        devis = Devis.objects.create(company=self.company, lead=self.lead,
                                     client=self.client_a,
                                     reference='DEV-ACAL295-1')
        Calepinage.objects.filter(pk=self.k.pk).update(devis=devis)
        avant = Calepinage.objects.count()
        reponse = self.api_u.post(f'{URL}depuis-modele/',
                                  {'devis_id': devis.pk}, format='json')
        self.assertEqual(reponse.status_code, 409, reponse.data)
        self.assertEqual(reponse.data, {'detail': MESSAGE_HORS_VUE})
        self.assertEqual(Calepinage.objects.count(), avant)

    def test_approbateur_voit_tout(self):
        reponse = self.api.post(f'{URL}comparer-projets/',
                                {'ids': [self.k.pk, self.m.pk]},
                                format='json')
        self.assertEqual(reponse.status_code, 200, reponse.data)
        self.assertEqual(len(reponse.data['lignes']), 2)
        self.assertEqual(reponse.data['refus'], [])
        reponse = self.api.post(f'{URL}depuis-lead/',
                                {'lead': self.lead_2.pk}, format='json')
        self.assertEqual(reponse.status_code, 200, reponse.data)
        self.assertEqual(reponse.data['calepinage'], self.m.pk)
