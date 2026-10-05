"""ACAL81 — le système de fixation CHOISI est persisté sur le calepinage.

Avant : rien ne gardait le choix — avec deux systèmes actifs, la
nomenclature et la feuille « Fixation » du classeur répondaient « Plusieurs
systèmes de fixation sont actifs… » à chaque rechargement (porte ATL-11).

Services réels (``services/fixation.py``, ``services/export_tableur.py``),
base réelle : aucun mock.
"""
from __future__ import annotations

import io

from apps.calepinage.models import (
    Calepinage, ComposantFixation, SystemeFixation,
)

from .test_api_liste import BaseApiCalepinage, url_detail
from .test_calx359_bom_fixation import deux_rangees_de_cinq, layout


def url_fixation(pk):
    return f'{url_detail(pk)}fixation/'


def url_bom(pk):
    return f'{url_detail(pk)}bom-fixation/'


class ChoixDuSystemeTest(BaseApiCalepinage):

    def setUp(self):
        super().setUp()
        self.calepinage = Calepinage.objects.create(
            company=self.company, lead_id=self.lead.pk, titre='Villa Anfa',
            roof_layout=layout(deux_rangees_de_cinq()))
        self.s1 = self._systeme('s1')
        self.s2 = self._systeme('s2')

    def _systeme(self, code, company=None):
        company = company or self.company
        systeme = SystemeFixation.objects.create(
            company=company, code=code, libelle=f'Système {code}',
            provenance='Notice fabricant (essai)')
        ComposantFixation.objects.create(
            company=company, systeme=systeme, role='rail',
            libelle=f'Rail {code}', unite='m', source='Notice', ordre=1,
            regle={'base': 'longueur_rangees_m', 'facteur': 2})
        return systeme

    def test_choix_persiste_lu_par_bom_et_classeur(self):
        from openpyxl import load_workbook

        from apps.calepinage.services.export_tableur import exporter_xlsx

        avant = self.api.get(url_bom(self.calepinage.pk))
        self.assertIsNone(avant.data['systeme'], 'deux actifs, aucun choix')

        reponse = self.api.post(url_fixation(self.calepinage.pk),
                                {'systeme_id': self.s2.pk}, format='json')
        self.assertEqual(reponse.status_code, 200, reponse.data)
        self.assertEqual(reponse.data['systeme']['id'], self.s2.pk)
        self.assertEqual(reponse.data['systeme_source'], 'calepinage')
        self.calepinage.refresh_from_db()
        self.assertEqual(self.calepinage.systeme_fixation_id, self.s2.pk)

        bom = self.api.get(url_bom(self.calepinage.pk))
        self.assertEqual(bom.status_code, 200, bom.data)
        self.assertEqual(bom.data['systeme']['id'], self.s2.pk)
        self.assertEqual(bom.data['systeme_source'], 'calepinage')
        self.assertEqual(bom.data['refus'], [])

        classeur = load_workbook(io.BytesIO(exporter_xlsx(
            Calepinage.objects.get(pk=self.calepinage.pk))))
        feuille = classeur['Fixation']
        textes = [str(cellule.value or '') for ligne in feuille.iter_rows()
                  for cellule in ligne]
        self.assertTrue(any('Rail s2' in texte for texte in textes))
        self.assertFalse(any('Plusieurs systèmes' in texte
                             for texte in textes))

    def test_autre_societe_et_absent_meme_refus(self):
        etranger = self._systeme('etranger', company=self.autre)
        refus = []
        for identifiant in (etranger.pk, 999999):
            reponse = self.api.post(url_fixation(self.calepinage.pk),
                                    {'systeme_id': identifiant},
                                    format='json')
            self.assertEqual(reponse.status_code, 400, reponse.data)
            self.assertIn('systeme_id', reponse.data)
            refus.append(str(reponse.data['systeme_id'])
                         .replace(str(identifiant), '<id>'))
        self.assertEqual(refus[0], refus[1])
        self.calepinage.refresh_from_db()
        self.assertIsNone(self.calepinage.systeme_fixation_id)

    def test_null_efface(self):
        self.api.post(url_fixation(self.calepinage.pk),
                      {'systeme_id': self.s1.pk}, format='json')
        reponse = self.api.post(url_fixation(self.calepinage.pk),
                                {'systeme_id': None}, format='json')
        self.assertEqual(reponse.status_code, 200, reponse.data)
        self.assertEqual(reponse.data,
                         {'systeme': None, 'systeme_source': None})
        self.calepinage.refresh_from_db()
        self.assertIsNone(self.calepinage.systeme_fixation_id)
        bom = self.api.get(url_bom(self.calepinage.pk))
        self.assertIsNone(bom.data['systeme'])

    def test_parametre_explicite_prime(self):
        self.api.post(url_fixation(self.calepinage.pk),
                      {'systeme_id': self.s2.pk}, format='json')
        bom = self.api.get(f'{url_bom(self.calepinage.pk)}'
                           f'?systeme={self.s1.pk}')
        self.assertEqual(bom.data['systeme']['id'], self.s1.pk)
        self.assertEqual(bom.data['systeme_source'], 'parametre')

    def test_calepinage_d_une_autre_societe_introuvable(self):
        reponse = self.api_autre.post(url_fixation(self.calepinage.pk),
                                      {'systeme_id': self.s1.pk},
                                      format='json')
        self.assertEqual(reponse.status_code, 404)
