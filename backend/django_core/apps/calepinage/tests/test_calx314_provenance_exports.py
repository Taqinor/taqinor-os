"""CALX314 — la provenance posée sur TOUS les exports, par UNE fonction.

Ce qui est prouvé ici :

* UNE seule fonction (``provenance_document.lignes_de_provenance``) produit
  les trois blocs : la feuille ``Provenance`` du XLSX, le texte du calque DXF
  ``PROVENANCE`` et le bloc ``provenance`` de l'export JSON (CALX312) — un
  test SONDE la fonction et retrouve sa sonde dans les trois sorties, un autre
  compare les trois sorties ligne à ligne ;
* elle REPREND la composition du CSV de simulation (``export_csv``) : ses six
  premières lignes coïncident avec l'en-tête de provenance du CSV ;
* une empreinte absente écrit « non calculée », jamais une chaîne vide ;
* la feuille ``Provenance`` est EN TÊTE du classeur ; le calque
  ``PROVENANCE`` est NON IMPRIMABLE ; le DXF se relit par ``ezdxf`` ;
* sans le paramètre, ``classeur_octets`` et ``document_dxf`` rendent
  exactement CAL179 / CAL178 (leurs tests restent verts).

Essais PURS : ni base, ni WeasyPrint.

Run :
    python manage.py test apps.calepinage.tests.test_calx314_provenance_exports
"""
import copy
import io
import json
import pathlib
import unittest
from types import SimpleNamespace
from unittest import mock

from apps.calepinage.services import provenance_document
from apps.calepinage.services.export_dxf import (
    CALQUE_MODULES, CALQUE_PROVENANCE, exporter_dxf, octets_dxf,
)
from apps.calepinage.services.export_projet import document_de_projet
from apps.calepinage.services.export_tableur import (
    FEUILLES, classeur_octets, exporter_xlsx, tables_du_resultat,
)
from apps.calepinage.services.planche import geometrie_de_planche
from apps.calepinage.services.provenance_document import (
    LIBELLE_EMPREINTE_LAYOUT, LIBELLE_EMPREINTE_SIMULATION, NON_CALCULEE,
    TITRE_FEUILLE, lignes_de_provenance,
)

from .test_cal171_planche import LAYOUT

ECHANTILLONS = pathlib.Path(__file__).resolve().parents[1] / 'contract_samples'
RESULTAT = json.loads((ECHANTILLONS / 'calepinage_resultat.json')
                      .read_text(encoding='utf-8'))['exemple']

#: Ce que la simulation a DÉPOSÉ sur le calepinage (``Calepinage.resultat``).
STOCKE = {
    'production': copy.deepcopy(RESULTAT['production']),
    'pertes': copy.deepcopy(RESULTAT['pertes']),
    'simulation': {'hash_entree': 'cd' * 32,
                   'calcule_le': '2026-09-21T10:15:00Z'},
}


def calepinage(**champs):
    """Un calepinage NON ENREGISTRÉ, sans société : aucune lecture en base."""
    valeurs = dict(pk=7, company=None, client_id=None, lead_id=None,
                   devis_id=None, titre='Villa Anfa',
                   roof_layout=copy.deepcopy(LAYOUT), layout_hash='ab' * 32,
                   version_moteur='calepinage-1.0.0',
                   resultat=copy.deepcopy(STOCKE))
    valeurs.update(champs)
    return SimpleNamespace(**valeurs)


def relire_xlsx(octets):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(octets))


def relire_dxf(octets):
    import ezdxf

    return ezdxf.read(io.StringIO(octets.decode('utf-8')))


def lignes_xlsx(octets):
    feuille = relire_xlsx(octets)[TITRE_FEUILLE]
    rangees = list(feuille.iter_rows(values_only=True))
    return [tuple(r) for r in rangees[1:]]


def lignes_dxf(octets):
    document = relire_dxf(octets)
    blocs = [e for e in document.modelspace()
             if e.dxf.layer == CALQUE_PROVENANCE]
    assert len(blocs) == 1, blocs
    return [tuple(ligne.split(' : ', 1))
            for ligne in blocs[0].plain_text().splitlines()]


def lignes_du_json(objet):
    document = document_de_projet(objet, resultat=None, site={},
                                  equipements={})
    return [(ligne['libelle'], ligne['valeur'])
            for ligne in document['provenance']]


class UneSeuleFonctionTest(unittest.TestCase):
    def test_une_sonde_dans_la_fonction_ressort_dans_les_trois_blocs(self):
        sonde = [('Sonde CALX314', 'valeur-sonde')]
        objet = calepinage()
        with mock.patch.object(provenance_document, 'lignes_de_provenance',
                               return_value=sonde) as fonction:
            xlsx = lignes_xlsx(exporter_xlsx(objet))
            dxf = lignes_dxf(exporter_dxf(objet))
            json_ = lignes_du_json(objet)
        self.assertEqual(fonction.call_count, 3)
        self.assertEqual(xlsx, sonde)
        self.assertEqual(dxf, sonde)
        self.assertEqual(json_, sonde)

    def test_les_trois_sorties_coincident_ligne_a_ligne(self):
        objet = calepinage()
        attendues = lignes_de_provenance(objet)
        self.assertEqual(len(attendues), 8)
        self.assertEqual(lignes_xlsx(exporter_xlsx(objet)), attendues)
        self.assertEqual(lignes_dxf(exporter_dxf(objet)), attendues)
        self.assertEqual(lignes_du_json(objet), attendues)

    def test_la_composition_reprend_celle_du_csv(self):
        from apps.calepinage.services.export_csv import _lignes_de_provenance
        from apps.calepinage.views.export_csv import document_exportable

        objet = calepinage()
        csv = [tuple(ligne) for ligne in
               _lignes_de_provenance(document_exportable(objet)) if ligne]
        self.assertEqual(lignes_de_provenance(objet)[:len(csv)], csv)
        self.assertEqual(len(csv), 6)

    def test_les_empreintes_sont_celles_du_calepinage_et_de_la_simulation(
            self):
        lignes = dict(lignes_de_provenance(calepinage()))
        self.assertEqual(lignes[LIBELLE_EMPREINTE_LAYOUT], 'ab' * 32)
        self.assertEqual(lignes[LIBELLE_EMPREINTE_SIMULATION], 'cd' * 32)
        self.assertEqual(lignes['Version du moteur'], 'calepinage-1.0.0')


class EmpreinteAbsenteTest(unittest.TestCase):
    def setUp(self):
        self.objet = calepinage(layout_hash='', resultat=None,
                                version_moteur='')

    def test_une_empreinte_absente_ecrit_non_calculee(self):
        lignes = dict(lignes_de_provenance(self.objet))
        self.assertEqual(lignes[LIBELLE_EMPREINTE_LAYOUT], NON_CALCULEE)
        self.assertEqual(lignes[LIBELLE_EMPREINTE_SIMULATION], NON_CALCULEE)
        self.assertEqual(NON_CALCULEE, 'non calculée')

    def test_aucune_valeur_n_est_une_chaine_vide(self):
        for libelle, valeur in lignes_de_provenance(self.objet):
            self.assertTrue(str(valeur).strip(), libelle)
        for libelle, valeur in lignes_xlsx(exporter_xlsx(self.objet)):
            self.assertTrue(str(valeur or '').strip(), libelle)

    def test_le_dxf_ecrit_non_calculee(self):
        lignes = dict(lignes_dxf(exporter_dxf(self.objet)))
        self.assertEqual(lignes[LIBELLE_EMPREINTE_LAYOUT], NON_CALCULEE)


class FormeDesSortiesTest(unittest.TestCase):
    def test_la_feuille_provenance_est_en_tete_du_classeur(self):
        classeur = relire_xlsx(exporter_xlsx(calepinage()))
        self.assertEqual(classeur.sheetnames,
                         [TITRE_FEUILLE] + list(FEUILLES))

    def test_le_calque_provenance_est_non_imprimable_et_le_dxf_se_relit(self):
        document = relire_dxf(exporter_dxf(calepinage()))
        self.assertIn(CALQUE_PROVENANCE, document.layers)
        self.assertEqual(document.layers.get(CALQUE_PROVENANCE).dxf.plot, 0)
        modules = [e for e in document.modelspace()
                   if e.dxf.layer == CALQUE_MODULES]
        self.assertEqual(len(modules), 2)

    def test_le_plan_de_cablage_porte_aussi_sa_provenance(self):
        from apps.calepinage.services.documents.plan_cablage import (
            exporter_plan_cablage_dxf,
        )

        affectation = [{'module': 'Pan Sud#%d' % rang, 'pan': 'Pan Sud',
                        'chaine': 1, 'onduleur': 1, 'mppt': 1,
                        'source': 'automatique'} for rang in (1, 2)]
        objet = calepinage()
        octets = exporter_plan_cablage_dxf(objet, affectation=affectation)
        self.assertEqual(lignes_dxf(octets), lignes_de_provenance(objet))

    def test_sans_le_parametre_cal178_et_cal179_sont_inchanges(self):
        geometrie = geometrie_de_planche(LAYOUT)
        document = relire_dxf(octets_dxf(geometrie))
        self.assertNotIn(CALQUE_PROVENANCE, document.layers)
        classeur = relire_xlsx(classeur_octets(
            tables_du_resultat(geometrie, STOCKE)))
        self.assertEqual(classeur.sheetnames, list(FEUILLES))

    def test_aucun_montant_dans_la_provenance(self):
        from apps.calepinage.services.export_tableur import (
            MOTS_D_ARGENT, verifier_absence_de_prix,
        )

        lignes = [list(ligne) for ligne in lignes_de_provenance(calepinage())]
        verifier_absence_de_prix(['Grandeur', 'Valeur'], lignes)  # ne lève pas
        texte = ' '.join(' '.join(ligne) for ligne in lignes).lower()
        for mot in MOTS_D_ARGENT:
            self.assertNotIn(mot, texte)


if __name__ == '__main__':  # pragma: no cover
    unittest.main()
