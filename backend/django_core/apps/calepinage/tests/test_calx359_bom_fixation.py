"""CALX359 — la nomenclature de fixation d'un calepinage, et son export.

Ce qui est prouvé ici :

* les GRANDEURS sont lues sur le document : modules, rangées, pans,
  jonctions entre voisins, extrémités (2 par segment continu), longueur des
  rangées (étendue + pas RELEVÉ) — un trou coupe la rangée, un pan tourné
  groupe ses rangées selon son azimut (plein sud : exactement le groupement
  de ``export_tableur.rangees_du_pan``), un pan sans voisin ne publie pas de
  longueur et le NOMME ;
* le cœur PUR ``_lignes_de_fixation`` rend EXACTEMENT les lignes de
  l'exemple committé (contrat CALX335) ; une règle non saisie ou un
  paramètre déclaré à ``null`` sortent à ``quantite: None`` avec ``manquant``
  nommé, jamais une quantité de repli ; une quantité « u » s'arrondit à
  l'unité supérieure ; une référence circulaire ou vers un rôle absent est
  dite ;
* catalogue vide ⇒ la réponse EST ``exemple_vide`` ;
* aucune valeur normative dans ``services/fixation.py`` : un test de surface
  refuse toute constante numérique hors 0 / 1 / 2 ;
* en HTTP : la porte sert la sortie du service, borne le système à la
  société, et la feuille « Fixation » n'apparaît au classeur QUE si la
  société a un catalogue (D12).

Les classes ``…EnBase`` exigent l'ORM : la CI est leur gate.

Run :
    python manage.py test apps.calepinage.tests.test_calx359_bom_fixation -v2
"""
import ast
import copy
import json
import math
import pathlib
from types import SimpleNamespace

from django.test import SimpleTestCase

from apps.calepinage.services import fixation
from apps.calepinage.services.export_tableur import rangees_du_pan

RACINE_APP = pathlib.Path(__file__).resolve().parents[1]
CONTRAT = json.loads(
    (RACINE_APP / 'contract_samples' / 'calepinage_fixation_bom.json')
    .read_text(encoding='utf-8'))
PAS = 1.2


def layout(panneaux, azimut=180.0):
    """Un document minimal : un pan, ses centres posés (mètres, ENU)."""
    return {
        'version': 2,
        'outline': [[33.5, -7.6], [33.5, -7.5998], [33.5002, -7.5998],
                    [33.5002, -7.6]],
        'zones': [{
            'id': 'z1', 'label': 'Pan Sud',
            'vertices': [[-7.6, 33.5], [-7.5998, 33.5], [-7.5998, 33.5002],
                         [-7.6, 33.5002]],
            'geometry': {
                'azimuthDeg': azimut, 'tiltDeg': 15.0,
                'count': len(panneaux), 'origin': [-7.6, 33.5],
                'panels': [{'cx': x, 'cy': y} for x, y in panneaux],
            },
        }],
    }


def deux_rangees_de_cinq():
    return [(1.0 + PAS * k, y) for y in (1.0, 3.0) for k in range(5)]


def composant(libelle, role, unite, regle, produit_id=None):
    return SimpleNamespace(libelle=libelle, role=role, unite=unite,
                           regle=regle, produit_id=produit_id)


def composants_du_contrat():
    return [
        composant('Rail aluminium', 'rail', 'm',
                  {'base': 'longueur_rangees_m', 'facteur': 2,
                   'libelle_facteur': 'rails par rangée'}, 812),
        composant('Pince de milieu', 'pince_milieu', 'u',
                  {'base': 'jonctions', 'facteur': 2,
                   'libelle_facteur': 'rails par rangée'}, 813),
        composant('Pince de fin', 'pince_fin', 'u',
                  {'base': 'extremites', 'facteur': 2,
                   'libelle_facteur': 'rails par rangée'}, 814),
        composant('Crochet de toit', 'crochet', 'u',
                  {'base': 'role:rail', 'diviseur': None,
                   'libelle_diviseur': 'entraxe des crochets (m)'}, 815),
        composant('Embout de rail', 'embout', 'u', {}),
    ]


def grandeurs(**valeurs):
    base = {'modules': 10, 'rangees': 2, 'pans': 1, 'jonctions': 8,
            'extremites': 4, 'longueur_rangees_m': 13.2}
    base.update(valeurs)
    return {'valeurs': base, 'manquants': {}}


class SurfaceSansValeurNormativeTest(SimpleTestCase):

    def test_aucune_constante_hors_0_1_2(self):
        texte = (RACINE_APP / 'services' / 'fixation.py').read_text(
            encoding='utf-8')
        intrus = [(noeud.lineno, noeud.value)
                  for noeud in ast.walk(ast.parse(texte))
                  if isinstance(noeud, ast.Constant)
                  and isinstance(noeud.value, (int, float))
                  and not isinstance(noeud.value, bool)
                  and noeud.value not in (0, 1, 2)]
        self.assertEqual(
            intrus, [],
            "Constante numérique dans services/fixation.py : un pas, un "
            "entraxe ou un nombre de rails se SAISIT sur le composant, il ne "
            "s'écrit jamais dans le code (CALX359).")


class GrandeursDuDocumentTest(SimpleTestCase):

    def test_deux_rangees_de_cinq_plein_sud(self):
        lues = fixation._grandeurs_du_document(layout(deux_rangees_de_cinq()))
        self.assertEqual(lues['manquants'], {})
        valeurs = lues['valeurs']
        self.assertEqual(valeurs['modules'], 10)
        self.assertEqual(valeurs['rangees'], 2)
        self.assertEqual(valeurs['pans'], 1)
        self.assertEqual(valeurs['jonctions'], 8)
        self.assertEqual(valeurs['extremites'], 4)
        self.assertAlmostEqual(valeurs['longueur_rangees_m'], 12.0, places=2)

    def test_un_trou_coupe_la_rangee(self):
        panneaux = [(0.0, 1.0), (PAS, 1.0), (PAS * 3, 1.0), (PAS * 4, 1.0)]
        valeurs = fixation._grandeurs_du_document(layout(panneaux))['valeurs']
        self.assertEqual(valeurs['rangees'], 1)
        self.assertEqual(valeurs['jonctions'], 2)
        self.assertEqual(valeurs['extremites'], 4)
        self.assertAlmostEqual(valeurs['longueur_rangees_m'], 4.8, places=2)

    def test_pan_tourne_groupe_selon_son_azimut(self):
        azimut = 150.0
        angle = math.radians(azimut)
        pente = (math.sin(angle), math.cos(angle))
        rangee = (math.cos(angle), -math.sin(angle))
        panneaux = [(5 + k * PAS * rangee[0] + r * 2 * pente[0],
                     5 + k * PAS * rangee[1] + r * 2 * pente[1])
                    for r in (0, 1) for k in range(5)]
        valeurs = fixation._grandeurs_du_document(
            layout(panneaux, azimut))['valeurs']
        self.assertEqual(valeurs['rangees'], 2)
        self.assertEqual(valeurs['jonctions'], 8)
        self.assertAlmostEqual(valeurs['longueur_rangees_m'], 12.0, places=2)

    def test_plein_sud_meme_groupement_que_le_tableur(self):
        modules = [(1.0, 1.0), (2.2, 1.0), (1.0, 3.0), (2.2, 3.004),
                   (3.4, 5.0)]
        rangs = rangees_du_pan(modules)
        attendu = sorted(sorted(round(c[0], 6) for c in modules
                                if rangs[c] == rang)
                         for rang in set(rangs.values()))
        # Plein sud, l'axe de position est l'est RETOURNÉ : on compare la
        # PARTITION des modules en rangées, pas le signe de l'axe.
        obtenu = sorted(sorted(round(-position, 6) for position in rangee)
                        for rangee in fixation._rangees_orientees(modules,
                                                                  180.0))
        self.assertEqual(obtenu, attendu)

    def test_pan_sans_voisin_nomme_la_longueur(self):
        panneaux = [(1.0, 1.0), (1.0, 3.0)]
        lues = fixation._grandeurs_du_document(layout(panneaux))
        self.assertIsNone(lues['valeurs']['longueur_rangees_m'])
        self.assertIn('Pan Sud', lues['manquants']['longueur_rangees_m'])
        self.assertEqual(lues['valeurs']['jonctions'], 0)
        self.assertEqual(lues['valeurs']['extremites'], 4)

    def test_sans_conception_ou_sans_module_tout_est_nomme(self):
        for document in (None, {}, layout([])):
            with self.subTest(document=bool(document)):
                lues = fixation._grandeurs_du_document(document)
                self.assertTrue(all(valeur is None for valeur
                                    in lues['valeurs'].values()))
                self.assertTrue(all(motif.strip() for motif
                                    in lues['manquants'].values()))


class LignesDeFixationTest(SimpleTestCase):

    def test_rejoue_exactement_l_exemple_committe(self):
        lignes = fixation._lignes_de_fixation(
            composants_du_contrat(), grandeurs(), {812, 813, 814, 815})
        self.assertEqual(lignes, CONTRAT['exemple']['lignes'])

    def test_diviseur_saisi_calcule_et_arrondit_a_l_unite_superieure(self):
        composants = composants_du_contrat()
        composants[3].regle = dict(composants[3].regle, diviseur=0.8)
        crochet = fixation._lignes_de_fixation(composants, grandeurs())[3]
        self.assertEqual(crochet['quantite'], 33)
        self.assertIsNone(crochet['manquant'])
        self.assertEqual(crochet['regle'],
                         'quantité de « rail » ÷ 0,8 (entraxe des crochets (m))')
        composants[3].regle = dict(composants[3].regle, diviseur=0.7)
        crochet = fixation._lignes_de_fixation(composants, grandeurs())[3]
        self.assertEqual(crochet['quantite'], math.ceil(26.4 / 0.7))

    def test_reference_calculee_meme_declaree_apres(self):
        composants = list(reversed(composants_du_contrat()))
        lignes = fixation._lignes_de_fixation(composants, grandeurs())
        self.assertEqual([ligne['role'] for ligne in lignes],
                         ['embout', 'crochet', 'pince_fin', 'pince_milieu',
                          'rail'])
        self.assertEqual(lignes[4]['quantite'], 26.4)

    def test_grandeur_illisible_propage_son_motif(self):
        lues = {'valeurs': dict(grandeurs()['valeurs'],
                                longueur_rangees_m=None),
                'manquants': {'longueur_rangees_m': 'pas illisible (essai)'}}
        lignes = fixation._lignes_de_fixation(composants_du_contrat(), lues)
        self.assertIsNone(lignes[0]['quantite'])
        self.assertEqual(lignes[0]['manquant'], 'pas illisible (essai)')
        self.assertIsNone(lignes[3]['quantite'])
        self.assertIn('« rail »', lignes[3]['manquant'])

    def test_reference_circulaire_ou_absente(self):
        boucle = [composant('A', 'rail', 'm', {'base': 'role:crochet'}),
                  composant('B', 'crochet', 'u', {'base': 'role:rail'}),
                  composant('C', 'visserie', 'u', {'base': 'role:lest'})]
        lignes = fixation._lignes_de_fixation(boucle, grandeurs())
        self.assertTrue(all(ligne['quantite'] is None for ligne in lignes))
        self.assertIn('circulaire', lignes[0]['manquant'])
        self.assertIn('circulaire', lignes[1]['manquant'])
        self.assertIn("qu'aucun composant", lignes[2]['manquant'])

    def test_produit_etranger_publie_null(self):
        lignes = fixation._lignes_de_fixation(composants_du_contrat(),
                                              grandeurs(), {813})
        self.assertIsNone(lignes[0]['produit_id'])
        self.assertEqual(lignes[1]['produit_id'], 813)

    def test_jamais_zero_de_repli(self):
        for ligne in fixation._lignes_de_fixation(composants_du_contrat(),
                                                  grandeurs()):
            self.assertNotEqual(ligne['quantite'], 0)


class CatalogueVideTest(SimpleTestCase):

    def test_la_reponse_est_l_exemple_vide(self):
        calepinage = SimpleNamespace(company=None, roof_layout=None)
        systeme, refus = fixation.resoudre_systeme(calepinage.company)
        self.assertIsNone(systeme)
        self.assertEqual(fixation.bom_de_fixation(calepinage, systeme, refus),
                         CONTRAT['exemple_vide'])

    def test_aucune_feuille_de_fixation_sans_catalogue(self):
        calepinage = SimpleNamespace(company=None, roof_layout=None)
        self.assertIsNone(fixation.table_fixation(calepinage))


# ── ORM — la CI est la gate de ces classes ─────────────────────────────────

import io  # noqa: E402

from apps.calepinage.models import (  # noqa: E402
    Calepinage, ComposantFixation, SystemeFixation,
)

from .test_api_liste import BaseApiCalepinage, url_detail  # noqa: E402


def url_bom(pk, **params):
    suite = ''.join(f'?{cle}={valeur}' for cle, valeur in params.items())
    return f'{url_detail(pk)}bom-fixation/{suite}'


class BomFixationEnBase(BaseApiCalepinage):

    def setUp(self):
        super().setUp()
        self.calepinage = Calepinage.objects.create(
            company=self.company, lead_id=self.lead.pk, titre='Villa Anfa',
            roof_layout=layout(deux_rangees_de_cinq()))

    def _systeme(self, company=None, code='incline'):
        company = company or self.company
        systeme = SystemeFixation.objects.create(
            company=company, code=code, libelle=f'Système {code}',
            provenance='Notice fabricant (essai)')
        ComposantFixation.objects.create(
            company=company, systeme=systeme, role='rail', libelle='Rail',
            unite='m', source='Notice', ordre=1,
            regle={'base': 'longueur_rangees_m', 'facteur': 2})
        ComposantFixation.objects.create(
            company=company, systeme=systeme, role='pince_milieu',
            libelle='Pince de milieu', unite='u', source='Notice', ordre=2,
            regle={'base': 'jonctions', 'facteur': 2})
        return systeme

    def test_catalogue_vide_sert_l_exemple_vide(self):
        reponse = self.api.get(url_bom(self.calepinage.pk))
        self.assertEqual(reponse.status_code, 200, reponse.data)
        self.assertEqual(reponse.data, CONTRAT['exemple_vide'])

    def test_unique_systeme_actif_applique(self):
        systeme = self._systeme()
        reponse = self.api.get(url_bom(self.calepinage.pk))
        self.assertEqual(reponse.status_code, 200, reponse.data)
        self.assertEqual(reponse.data['systeme']['id'], systeme.pk)
        self.assertEqual([ligne['quantite'] for ligne in reponse.data['lignes']],
                         [24.0, 16])
        self.assertEqual(sorted(reponse.data), ['lignes', 'refus', 'systeme'])

    def test_plusieurs_systemes_sans_choix_refus_nomme(self):
        self._systeme(code='a')
        choisi = self._systeme(code='b')
        reponse = self.api.get(url_bom(self.calepinage.pk))
        self.assertIsNone(reponse.data['systeme'])
        self.assertEqual(reponse.data['refus'][0]['champ'], 'systeme')
        reponse = self.api.get(url_bom(self.calepinage.pk,
                                       systeme=choisi.pk))
        self.assertEqual(reponse.data['systeme']['code'], 'b')

    def test_systeme_d_une_autre_societe_refuse(self):
        etranger = self._systeme(company=self.autre, code='etranger')
        reponse = self.api.get(url_bom(self.calepinage.pk,
                                       systeme=etranger.pk))
        self.assertIsNone(reponse.data['systeme'])
        self.assertEqual(reponse.data['refus'][0]['champ'], 'systeme')

    def test_calepinage_d_une_autre_societe_introuvable(self):
        reponse = self.api_autre.get(url_bom(self.calepinage.pk))
        self.assertEqual(reponse.status_code, 404)

    def test_feuille_fixation_au_classeur_seulement_avec_catalogue(self):
        from openpyxl import load_workbook

        from apps.calepinage.services.export_tableur import exporter_xlsx

        sans = load_workbook(io.BytesIO(exporter_xlsx(self.calepinage)))
        self.assertNotIn('Fixation', sans.sheetnames)
        self._systeme()
        avec = load_workbook(io.BytesIO(exporter_xlsx(self.calepinage)))
        self.assertEqual(avec.sheetnames[-1], 'Fixation')

    def test_rien_n_est_ecrit(self):
        self._systeme()
        avant = copy.deepcopy(self.calepinage.roof_layout)
        self.api.get(url_bom(self.calepinage.pk))
        self.calepinage.refresh_from_db()
        self.assertEqual(self.calepinage.roof_layout, avant)
