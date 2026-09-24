"""CALX341 — comparer jusqu'à 5 calepinages, et le comparatif en classeur.

CE QUE CE FICHIER PROUVE
------------------------
Deux étages, deux coûts :

1. Sans base (``SimpleTestCase``) — les deux actions sont RÉELLEMENT
   rattachées au viewset pivot (piège CALX7 : un ``__name__`` divergent fait
   disparaître la route), la lecture des ``ids`` refuse en NOMMANT le champ
   (six identifiants ⇒ la borne est dite), une ligne non simulée ou périmée
   porte ``null`` et jamais ``0`` avec son motif, et l'exemple COMMITTÉ du
   contrat (``calepinage_comparaison_projets.json``, CALX331) est RECALCULÉ
   ligne pour ligne depuis des calepinages en mémoire — c'est lui que la
   moitié frontend importe. Le classeur « Comparatif » s'ouvre, porte les
   colonnes du contrat, laisse vide ce qui n'est pas simulé et ne porte aucun
   prix.
2. En base (``…EnBaseTest``) — la porte HTTP : 6 ids ⇒ 400 nommant la borne,
   un calepinage d'une autre société est IGNORÉ et listé dans ``refus``, un
   calepinage non simulé rend ses colonnes à ``null`` avec son motif, et
   ``comparatif.xlsx`` sert un classeur.

Run :
    python manage.py test apps.calepinage.tests.test_calx341_comparaison_projets -v2
"""
from __future__ import annotations

import io
import json
import pathlib

from django.test import SimpleTestCase

from apps.calepinage.models import Calepinage
from apps.calepinage.services.comparaison import CLES_PRODUCTION
from apps.calepinage.services.comparaison_projets import (
    BORNE_PROJETS, ComparaisonRefusee, MOTIF_INTROUVABLE, MOTIF_NON_SIMULE,
    MOTIF_PERIME, _colonnes, _ligne_de_comparaison, _lire_ids,
)
from apps.calepinage.services.export_tableur import (
    FEUILLE_COMPARATIF, MOTS_D_ARGENT, exporter_comparatif_xlsx,
)

from .test_api_liste import BaseApiCalepinage, url_detail

RACINE = pathlib.Path(__file__).resolve().parent.parent
ECHANTILLON = (RACINE / 'contract_samples'
               / 'calepinage_comparaison_projets.json')

URL_COMPARER = '/api/django/calepinage/calepinages/comparer-projets/'


def _contrat():
    return json.loads(ECHANTILLON.read_text(encoding='utf-8'))


def _actions():
    from apps.calepinage import urls  # noqa: F401
    from apps.calepinage.views.calepinages import CalepinageViewSet

    return CalepinageViewSet, {methode.__name__: methode
                               for methode
                               in CalepinageViewSet.get_extra_actions()}


class Faux:
    """Un calepinage EN MÉMOIRE : les seuls attributs que le service lit."""

    def __init__(self, pk, *, titre='', statut='brouillon', layout_hash='',
                 roof_layout=None, resultat=None):
        self.pk = pk
        self.titre = titre
        self.statut = statut
        self.layout_hash = layout_hash
        self.roof_layout = roof_layout
        self.resultat = resultat


#: Les trois calepinages dont l'exemple committé est la comparaison.
A = 'ab' * 32
C = 'cd' * 32
E = 'ef' * 32


def _faux_de_l_exemple():
    simule = Faux(
        1, titre="Calepinage d'essai A — toiture sud", statut='valide',
        layout_hash=A,
        resultat={
            'pose': {'total_modules': 12, 'kwc': 8.64},
            'production': {'total': {
                'kwc': 8.64, 'p50_kwh': 13000.0, 'p75_kwh': 12400.0,
                'p90_kwh': 11900.0, 'performance_ratio': 0.799,
                'specific_yield_kwh_kwc': 1504.6}},
            'autoconsommation': {'self_consumption_rate': 0.62},
        })
    jamais_simule = Faux(
        2, titre="Calepinage d'essai B — ombrière", layout_hash=C,
        roof_layout={'result': {'panels': 20, 'kwc': 14.4}})
    perime = Faux(
        3, titre="Calepinage d'essai C — toiture est-ouest", layout_hash=E,
        roof_layout={'result': {'panels': 16, 'kwc': 11.52}},
        resultat={
            # La simulation a été calculée sur UNE AUTRE empreinte.
            'layout_hash': A,
            'production': {'total': {'kwc': 99.0, 'p50_kwh': 50000.0}},
        })
    return simule, jamais_simule, perime


class RoutageTest(SimpleTestCase):
    """Sans route, l'écran de comparaison appellerait dans le vide."""

    def test_comparer_projets_est_une_action_de_liste_en_post(self):
        _viewset, actions = _actions()
        self.assertIn('comparer_projets', actions)
        action = actions['comparer_projets']
        self.assertFalse(action.detail)
        self.assertEqual(action.url_path, 'comparer-projets')
        self.assertEqual(set(action.mapping), {'post'})
        self.assertEqual(
            [garde.__name__ for garde in action.kwargs['permission_classes']],
            ['PeutVoirCalepinage'])

    def test_comparatif_xlsx_est_une_sous_ressource_en_get(self):
        _viewset, actions = _actions()
        self.assertIn('comparatif_xlsx', actions)
        action = actions['comparatif_xlsx']
        self.assertTrue(action.detail)
        self.assertEqual(action.url_path, 'comparatif.xlsx')
        self.assertEqual(set(action.mapping), {'get'})
        self.assertEqual(
            [garde.__name__ for garde in action.kwargs['permission_classes']],
            ['PeutVoirCalepinage'])

    def test_les_noms_d_attribut_egalent_les_noms_de_fonction(self):
        viewset, _ = _actions()
        for nom in ('comparer_projets', 'comparatif_xlsx'):
            self.assertEqual(getattr(viewset, nom).__name__, nom)

    def test_le_module_est_declare_dans_MODULES_RATTACHES(self):
        from apps.calepinage.views.rattachements import MODULES_RATTACHES

        self.assertIn('comparaison_projets', MODULES_RATTACHES)


class LireIdsTest(SimpleTestCase):
    """Les refus NOMMENT le champ ``ids`` — jamais un 400 générique."""

    def test_six_ids_refuses_en_nommant_la_borne(self):
        with self.assertRaises(ComparaisonRefusee) as refus:
            _lire_ids([1, 2, 3, 4, 5, 6])
        self.assertEqual(refus.exception.champ, 'ids')
        self.assertIn('Au plus 5 calepinages', str(refus.exception))
        self.assertIn('reçu : 6', str(refus.exception))

    def test_cinq_ids_acceptes_la_borne_est_celle_de_pvsol(self):
        self.assertEqual(BORNE_PROJETS, 5)
        self.assertEqual(_lire_ids([5, 4, 3, 2, 1]), [5, 4, 3, 2, 1])

    def test_un_id_repete_ne_compte_qu_une_fois(self):
        self.assertEqual(_lire_ids([2, 1, 2, 3, 4, 5, 1]), [2, 1, 3, 4, 5])

    def test_liste_absente_ou_vide_refusee(self):
        for brut in (None, [], '', '  '):
            with self.subTest(brut=brut), \
                    self.assertRaises(ComparaisonRefusee) as refus:
                _lire_ids(brut)
            self.assertEqual(refus.exception.champ, 'ids')

    def test_jeton_illisible_refuse_en_le_citant(self):
        for brut in (['abc'], [0], [-3], [True], [1.5]):
            with self.subTest(brut=brut), \
                    self.assertRaises(ComparaisonRefusee) as refus:
                _lire_ids(brut)
            self.assertEqual(refus.exception.champ, 'ids')
            self.assertIn('entiers positifs', str(refus.exception))

    def test_chaine_de_requete_et_liste_melangee(self):
        self.assertEqual(_lire_ids('3,1'), [3, 1])
        self.assertEqual(_lire_ids([7, '2,3']), [7, 2, 3])


class LigneDeComparaisonTest(SimpleTestCase):
    """Non simulé ⇒ ``null`` (jamais ``0``) ET un motif qui dit pourquoi."""

    def test_non_simule_rend_null_et_son_motif(self):
        ligne = _ligne_de_comparaison(Faux(4, titre='Nu'))
        self.assertFalse(ligne['simule'])
        self.assertEqual(ligne['motif'], MOTIF_NON_SIMULE)
        for cle in CLES_PRODUCTION:
            self.assertIsNone(ligne[cle], cle)
        self.assertIsNone(ligne['modules'])
        self.assertIsNone(ligne['kwc'])

    def test_perime_rend_null_et_son_motif_et_ignore_le_kwc_perime(self):
        _simule, _jamais, perime = _faux_de_l_exemple()
        ligne = _ligne_de_comparaison(perime)
        self.assertFalse(ligne['simule'])
        self.assertEqual(ligne['motif'], MOTIF_PERIME)
        for cle in CLES_PRODUCTION:
            self.assertIsNone(ligne[cle], cle)
        # Le kWc de la simulation PÉRIMÉE (99) n'est jamais repris : c'est
        # celui de la conception courante qui est publié.
        self.assertEqual(ligne['kwc'], 11.52)

    def test_simule_rend_les_grandeurs_lues(self):
        simule, _jamais, _perime = _faux_de_l_exemple()
        ligne = _ligne_de_comparaison(simule)
        self.assertTrue(ligne['simule'])
        self.assertEqual(ligne['motif'], '')
        self.assertEqual(ligne['p50_kwh'], 13000.0)
        self.assertEqual(ligne['self_consumption_rate'], 0.62)
        self.assertEqual((ligne['modules'], ligne['kwc']), (12, 8.64))

    def test_un_zero_mesure_reste_un_zero(self):
        """``None`` pour « non mesuré », mais un 0 MESURÉ n'est pas effacé."""
        ligne = _ligne_de_comparaison(Faux(5, resultat={
            'production': {'total': {'p50_kwh': 0.0}}}))
        self.assertTrue(ligne['simule'])
        self.assertEqual(ligne['p50_kwh'], 0.0)


class ContratCommitteTest(SimpleTestCase):
    """L'exemple COMMITTÉ est RECALCULÉ par le service, ligne pour ligne."""

    def test_le_chemin_declare_est_celui_de_l_action(self):
        self.assertEqual(_contrat()['endpoint'], 'POST ' + URL_COMPARER)

    def test_les_colonnes_sont_celles_du_service(self):
        for etat in ('exemple', 'exemple_vide'):
            self.assertEqual(_contrat()[etat]['colonnes'], _colonnes(), etat)

    def test_les_trois_lignes_de_l_exemple_sont_recalculees(self):
        attendues = _contrat()['exemple']['lignes']
        obtenues = [_ligne_de_comparaison(faux)
                    for faux in _faux_de_l_exemple()]
        self.assertEqual(obtenues, attendues)

    def test_chaque_ligne_porte_les_cles_de_production(self):
        for ligne in _contrat()['exemple']['lignes']:
            self.assertLessEqual(set(CLES_PRODUCTION), set(ligne))

    def test_le_refus_porte_le_motif_du_service(self):
        for etat in ('exemple', 'exemple_vide'):
            self.assertEqual(_contrat()[etat]['refus'],
                             [{'id': 9, 'motif': MOTIF_INTROUVABLE}], etat)

    def test_plus_de_cinq_ids_est_documente_et_nomme_ids(self):
        contrat = _contrat()
        self.assertEqual(list(contrat['refus_plus_de_cinq']), ['ids'])
        with self.assertRaises(ComparaisonRefusee) as refus:
            _lire_ids([1, 2, 3, 4, 5, 6])
        self.assertEqual(contrat['refus_plus_de_cinq']['ids'],
                         str(refus.exception))

    def test_les_refus_documentes_sont_les_messages_du_service(self):
        contrat = _contrat()
        for cle, brut in (('refus_sans_ids', []),
                          ('refus_ids_illisibles', ['abc'])):
            with self.subTest(refus=cle), \
                    self.assertRaises(ComparaisonRefusee) as refus:
                _lire_ids(brut)
            self.assertEqual(contrat[cle], {'ids': str(refus.exception)})


class ClasseurComparatifTest(SimpleTestCase):
    """UNE feuille « Comparatif », les colonnes du contrat, aucun prix."""

    def _classeur(self):
        from openpyxl import load_workbook

        exemple = _contrat()['exemple']
        return load_workbook(io.BytesIO(exporter_comparatif_xlsx(exemple)))

    def test_une_feuille_comparatif_aux_colonnes_du_contrat(self):
        classeur = self._classeur()
        self.assertEqual(classeur.sheetnames, [FEUILLE_COMPARATIF])
        entetes = [cellule.value for cellule in classeur.active[1]]
        self.assertEqual(entetes[:3], ['Calepinage', 'Statut', 'Simulé'])
        self.assertEqual(entetes[-1], 'Motif')
        self.assertIn('Production P50 (kWh/an)', entetes)
        self.assertEqual(len(entetes), 3 + len(_colonnes()) + 1)

    def test_non_simule_reste_une_cellule_vide_jamais_zero(self):
        feuille = self._classeur().active
        entetes = [cellule.value for cellule in feuille[1]]
        p50 = entetes.index('Production P50 (kWh/an)')
        ligne_b = [cellule.value for cellule in feuille[3]]
        self.assertIn(ligne_b[p50], (None, ''))
        self.assertEqual(ligne_b[-1], MOTIF_NON_SIMULE)

    def test_les_refus_sont_listes_en_fin_de_feuille(self):
        feuille = self._classeur().active
        derniere = [cellule.value for cellule in feuille[feuille.max_row]]
        self.assertEqual(derniere[0], 'Calepinage #9')
        self.assertEqual(derniere[-1], MOTIF_INTROUVABLE)

    def test_aucun_mot_d_argent_dans_les_entetes(self):
        entetes = [str(cellule.value).lower()
                   for cellule in self._classeur().active[1]]
        for mot in MOTS_D_ARGENT:
            for entete in entetes:
                self.assertNotIn(mot, entete)

    def test_un_titre_client_ne_fait_pas_refuser_l_export(self):
        """« Hammadi » contient « mad » : le TITRE saisi n'est pas un prix."""
        comparaison = {'colonnes': _colonnes(), 'refus': [], 'lignes': [
            _ligne_de_comparaison(Faux(6, titre='Villa Hammadi'))]}
        self.assertTrue(exporter_comparatif_xlsx(comparaison))


class ComparaisonProjetsApiEnBaseTest(BaseApiCalepinage):
    """La porte HTTP — exige la base (CI)."""

    def setUp(self):
        super().setUp()
        self.a = Calepinage.objects.create(
            company=self.company, lead_id=self.lead.pk, titre='Villa A',
            layout_hash=A, roof_layout={'result': {'panels': 12}},
            resultat={'production': {'total': {
                'kwc': 8.64, 'p50_kwh': 13000.0}}})
        self.b = Calepinage.objects.create(
            company=self.company, lead_id=self.lead.pk, titre='Villa B')
        self.etranger = Calepinage.objects.create(
            company=self.autre, lead_id=999, titre='Chez la voisine')

    def test_six_ids_refuses_400_en_nommant_la_borne(self):
        reponse = self.api.post(URL_COMPARER,
                                {'ids': [1, 2, 3, 4, 5, 6]},
                                format='json')
        self.assertEqual(reponse.status_code, 400, reponse.data)
        self.assertIn('ids', reponse.data)
        self.assertIn('Au plus 5', str(reponse.data['ids']))

    def test_autre_societe_ignoree_et_listee_dans_refus(self):
        reponse = self.api.post(
            URL_COMPARER,
            {'ids': [self.a.pk, self.etranger.pk]}, format='json')
        self.assertEqual(reponse.status_code, 200, reponse.data)
        self.assertEqual([ligne['id'] for ligne in reponse.data['lignes']],
                         [self.a.pk])
        self.assertEqual(reponse.data['refus'],
                         [{'id': self.etranger.pk,
                           'motif': MOTIF_INTROUVABLE}])

    def test_non_simule_colonnes_null_et_motif(self):
        reponse = self.api.post(
            URL_COMPARER, {'ids': [self.b.pk, self.a.pk]}, format='json')
        self.assertEqual(reponse.status_code, 200, reponse.data)
        b, a = reponse.data['lignes']
        self.assertEqual((b['id'], a['id']), (self.b.pk, self.a.pk))
        self.assertFalse(b['simule'])
        self.assertEqual(b['motif'], MOTIF_NON_SIMULE)
        for cle in CLES_PRODUCTION:
            self.assertIsNone(b[cle], cle)
        self.assertTrue(a['simule'])
        self.assertEqual(a['p50_kwh'], 13000.0)
        self.assertEqual(sorted(reponse.data),
                         sorted(_contrat()['exemple']))

    def test_aucune_ecriture(self):
        avant = Calepinage.objects.get(pk=self.b.pk).updated_at
        self.api.post(URL_COMPARER, {'ids': [self.b.pk]}, format='json')
        self.assertEqual(Calepinage.objects.get(pk=self.b.pk).updated_at,
                         avant)

    def test_comparatif_xlsx_sert_un_classeur(self):
        reponse = self.api.get(
            url_detail(self.a.pk) + 'comparatif.xlsx/',
            {'ids': f'{self.b.pk},{self.etranger.pk}'})
        self.assertEqual(reponse.status_code, 200)
        self.assertIn('spreadsheetml', reponse['Content-Type'])
        self.assertIn('comparatif-calepinage-', reponse[
            'Content-Disposition'])

    def test_comparatif_xlsx_d_une_autre_societe_introuvable(self):
        reponse = self.api.get(
            url_detail(self.etranger.pk) + 'comparatif.xlsx/')
        self.assertEqual(reponse.status_code, 404)

    def test_comparatif_xlsx_six_calepinages_refuse(self):
        reponse = self.api.get(url_detail(self.a.pk) + 'comparatif.xlsx/',
                               {'ids': '101,102,103,104,105'})
        self.assertEqual(reponse.status_code, 400)
        self.assertIn('ids', reponse.data)
